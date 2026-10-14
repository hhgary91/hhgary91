"""Standard Design Library register v3 (Singapore).
Registers: Drain, RW (RC/gravity), Embedded wall (CBP/secant/sheet pile/D-wall), Manhole, ERSS pit, Linkway (LTA SDRE Ch.22), Other.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule
import datetime, os

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "Design_Library_Register.xlsx")
HDR, FIRST, LAST = 4, 5, 304

F = "Arial"
f_title = Font(name=F, size=14, bold=True, color="1F3864")
f_sub = Font(name=F, size=9, italic=True, color="595959")
f_hdr = Font(name=F, size=9, bold=True, color="FFFFFF")
f_in = Font(name=F, size=9, color="0000FF")
f_calc = Font(name=F, size=9, color="000000")
f_body = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_sec = Font(name=F, size=11, bold=True, color="1F3864")
fill_hdr = {"in": PatternFill("solid", fgColor="1F3864"), "calc": PatternFill("solid", fgColor="7F7F7F"),
            "qty": PatternFill("solid", fgColor="375623"), "file": PatternFill("solid", fgColor="7030A0")}
fill_in = PatternFill("solid", fgColor="FFF2CC")
fill_calc = PatternFill("solid", fgColor="F2F2F2")
fill_ex = PatternFill("solid", fgColor="E2EFDA")
fill_red = PatternFill("solid", fgColor="F8CBAD")
fill_grn = PatternFill("solid", fgColor="C6EFCE")
fill_amb = PatternFill("solid", fgColor="FFE699")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")
topwrap = Alignment(wrap_text=True, vertical="top")
TODAY = datetime.date(2026, 10, 14)

wb = Workbook()
ws_readme = wb.active
ws_readme.title = "README"


def put(ws, ref, v, font=f_body, fill=None, fmt=None, al=None):
    ws[ref] = v
    ws[ref].font = font
    ws[ref].border = box
    if fill:
        ws[ref].fill = fill
    if fmt:
        ws[ref].number_format = fmt
    if al:
        ws[ref].alignment = al


# =============================================================== SG_Criteria
st = wb.create_sheet("SG_Criteria")
st["A1"] = "Singapore Design Criteria & Library Settings"
st["A1"].font = f_title
st["A2"] = ("Compliance checks in every register read these cells. Rows marked VERIFY came from older editions or secondary sources: "
            "confirm against the current document before go-live and fill 'Verified by'.")
st["A2"].font = f_sub
for c, h in enumerate(["Parameter", "Value", "Unit", "Source / clause", "Status", "Verified by", "Date"], 1):
    put(st, f"{L(c)}3", h, f_hdr, fill_hdr["in"], al=wrap)
S = {}
srows = [
    ("sec", "LIBRARY FOLDERS (SharePoint document library synced to PC)"),
    ("root", "Library root folder", r"C:\Users\yourname\YourCompany\Design Library - Documents", "", "Path of the synced SharePoint library on your PC", "SET"),
    ("rootn", "Library root (normalised)", '=IF(RIGHT(B5,1)="\\",B5,B5&"\\")', "", "Automatic", "calc"),
    ("f_dr", "Drainage folder", "01_Drainage", "", "", "SET"),
    ("f_rw", "Retaining wall folder", "02_Retaining_Wall", "", "", "SET"),
    ("f_ew", "Embedded wall / ERSS folder", "03_Embedded_Wall", "", "", "SET"),
    ("f_mh", "Sewer manhole folder", "04_Sewer_Manhole", "", "", "SET"),
    ("f_ep", "Temporary ERSS pit folder", "05_ERSS_Pit", "", "", "SET"),
    ("f_lk", "Linkway folder", "06_Linkway", "", "", "SET"),
    ("f_ot", "Other structures folder", "07_Other_Structures", "", "", "SET"),
    ("sec", "DRAINAGE: PUB COP on Surface Water Drainage / LTA SDRE Ch.2"),
    ("rdW", "Roadside drain: min internal width", 600, "mm", "PUB COP SWD (roadside drain incl. entrance culvert)", "VERIFY"),
    ("rdD", "Roadside drain: min clear depth", 600, "mm", "PUB COP SWD (roadside drain incl. entrance culvert)", "VERIFY"),
    ("cdW", "Common drain: min internal width", 450, "mm", "PUB COP SWD cl. 4.5.1", "VERIFY"),
    ("rail", "Open drain depth above which safety railing is required", 1000, "mm", "PUB COP SWD safety railings clause", "VERIFY"),
    ("fck", "Min concrete grade, drains / walls (fck)", 28, "MPa", "LTA SDRE Ch.2 general notes (C28/35)", "VERIFY"),
    ("cov", "Min clear cover, drains / walls", 40, "mm", "LTA SDRE Ch.2 general notes", "VERIFY"),
    ("q12", "Design surcharge on drain walls", 12, "kPa", "LTA SDRE Ch.2 precast U-drain notes", "VERIFY"),
    ("q5", "Load on top slab of slabbed-over drain (footway)", 5, "kPa", "LTA SDRE Ch.2", "VERIFY"),
    ("vmin", "Min self-cleansing velocity", 1.0, "m/s", "PUB COP SWD", "VERIFY"),
    ("vmax", "Max velocity, concrete-lined drain", 3.0, "m/s", "PUB COP SWD", "VERIFY"),
    ("fb", "Freeboard used for capacity (fraction of depth)", 0.15, "", "Company assumption. Check PUB COP SWD cl. 7.3.4", "VERIFY"),
    ("nman", "Manning's n, concrete", 0.015, "", "Company assumption", "SET"),
    ("rp1", "Return period, catchment < 100 ha", 10, "years", "PUB COP SWD", "VERIFY"),
    ("rp2", "Return period, catchment 100 - 1000 ha", 25, "years", "PUB COP SWD", "VERIFY"),
    ("rp3", "Return period, catchment > 1000 ha", 50, "years", "PUB COP SWD", "VERIFY"),
    ("sec", "RETAINING STRUCTURES & ERSS: LTA CDC / BCA ERSS submission requirements (Mar 2024)"),
    ("qrw", "Min surcharge, permanent earth-retaining structure (LTA works)", 25, "kPa", "Older LTA CDC value. Confirm vs CDC C&S Rev A3 (Dec 2025)", "VERIFY"),
    ("life", "Design service life, permanent structures (LTA)", 120, "years", "LTA CDC durability clause", "VERIFY"),
    ("bca0", "ERSS: no BCA submission up to (RC / steel wall)", 1.5, "m", "BCA ERSS submission requirements, Mar 2024", "VERIFY"),
    ("bca0n", "ERSS: no BCA submission up to (other wall types)", 1.0, "m", "BCA ERSS submission requirements, Mar 2024", "VERIFY"),
    ("bca1", "ERSS: QP(ST) only up to", 4.0, "m", "BCA: 1.5-4 m QP(ST), no AC", "VERIFY"),
    ("bca2", "ERSS: QP(ST) + AC up to", 6.0, "m", "BCA: 4-6 m QP(ST) + AC; > 6 m geotechnical building works", "VERIFY"),
    ("fosb", "Min base stability FOS used for library ERSS pits", 1.5, "", "Company criterion", "SET"),
    ("sec", "SEWERAGE: PUB COP on Sewerage & Sanitary Works 3rd Ed (Mar 2025) & standard drawings"),
    ("mhd", "Min depth of sewers and manholes", 1.5, "m", "PUB COP SSW (2nd Ed wording; confirm in 3rd Ed)", "VERIFY"),
    ("sdia", "Min public sewer diameter", 250, "mm", "PUB COP SSW", "VERIFY"),
    ("mhs", "Max distance between manholes (project check)", 120, "m", "PUB COP SSW (2nd Ed; confirm in 3rd Ed)", "VERIFY"),
    ("mh6", "Max depth for standard manhole drawing PUB/WRN/STD/001C", 6.0, "m", "PUB standard drawings list (001C <= 6 m; 006C precast > 6 m)", "VERIFY"),
    ("mhbig", "Pipe dia from which rectangular manhole 004C applies", 1000, "mm", "PUB/WRN/STD/004C", "VERIFY"),
    ("upl", "Min uplift (flotation) FOS for manholes", 1.1, "", "Company criterion; check vs SS EN 1997-1 UPL", "SET"),
    ("sec", "COVERED LINKWAY: LTA SDRE Ch.22 (Rev I, Mar 2026) / LTA Architectural Design Criteria"),
    ("lkH", "Low covered linkway: min clear headroom", 2.4, "m", "LTA covered linkway guide / ADC s4.8", "VERIFY"),
    ("lkHH", "High covered linkway: min clear headroom", 5.7, "m", "LTA covered linkway guide", "VERIFY"),
    ("lkW", "Min roof width (eave to eave)", 2.4, "m", "LTA ADC s4.8", "VERIFY"),
    ("lkW2", "Min roof width to MRT/LRT or high-traffic routes", 3.6, "m", "LTA ADC s4.8", "VERIFY"),
    ("lkV", "Basic wind speed for linkway and fixings", 22, "m/s", "LTA SDRE Ch.22 notes; SS EN 1991-1-4", "VERIFY"),
    ("lkP", "Roof fixing ultimate pull-out capacity", 4.0, "kN", "LTA SDRE Ch.22 notes", "VERIFY"),
    ("lkC", "Min RC grade for linkway stumps / footings", 25, "MPa", "LTA SDRE Ch.22 notes (C25/30)", "VERIFY"),
]
r = 4
for row in srows:
    if row[0] == "sec":
        st[f"A{r}"] = row[1]
        st[f"A{r}"].font = f_sec
        r += 1
        continue
    key, lab, val, unit, src, status = row
    S[key] = f"SG_Criteria!$B${r}"
    put(st, f"A{r}", lab, f_bold)
    put(st, f"B{r}", val, f_calc if status == "calc" else f_in, fill_calc if status == "calc" else fill_in)
    put(st, f"C{r}", unit)
    put(st, f"D{r}", src, al=topwrap)
    put(st, f"E{r}", status, f_bold, fill_amb if status == "VERIFY" else None)
    put(st, f"F{r}", "", f_in, fill_in)
    put(st, f"G{r}", "", f_in, fill_in, "dd-mmm-yy")
    r += 1
assert S["root"] == "SG_Criteria!$B$5"
r += 1
st[f"A{r}"] = "REFERENCE DOCUMENTS (edition the library is built on)"
st[f"A{r}"].font = f_sec
r += 1
for c, h in enumerate(["Document", "Edition / revision", "", "Applies to", "Status", "Verified by", "Date"], 1):
    put(st, f"{L(c)}{r}", h, f_hdr, fill_hdr["in"])
r += 1
docs = [
    ("PUB Code of Practice on Surface Water Drainage", "7th Ed. Dec 2018, Addendum No. 3 (Apr 2025)", "Drains, culverts, sumps"),
    ("PUB Code of Practice on Sewerage and Sanitary Works", "3rd Ed. (Mar 2025)", "Sewers, manholes"),
    ("PUB Standard Drawings for Sewerage Works", "PUB/WRN/STD/001C, 003B, 004C, 006C", "Manholes <= 6 m, precast parts, rectangular, > 6 m"),
    ("LTA Standard Details of Road Elements (SDRE)", "Apr 2014 Ed., Revision I (Mar 2026)", "Ch.2 Drains & Culverts; Ch.22 Covered Linkway (CLW1-CLW5)"),
    ("LTA Architectural Design Criteria s4.8 / Covered linkway guide", "Current", "Linkway headroom, width"),
    ("LTA Civil Design Criteria for Civil & Structural Works", "Revision A3 (Dec 2025)", "Loads, surcharge, durability for LTA works"),
    ("BCA ERSS submission requirements", "Mar 2024", "Submission thresholds, QP / AC"),
    ("BCA guideline on integrity of sheet pile walls", "Nov 2025", "Sheet pile ERSS"),
    ("SS EN 1992-1-1, SS EN 1991-2, SS EN 1991-1-4, SS EN 1997-1 + Singapore NAs", "Current", "RC, traffic, wind, geotechnical"),
]
for d in docs:
    put(st, f"A{r}", d[0], f_bold)
    put(st, f"B{r}", d[1], f_in, fill_in)
    put(st, f"D{r}", d[2])
    put(st, f"E{r}", "VERIFY", f_bold, fill_amb)
    put(st, f"F{r}", "", f_in, fill_in)
    put(st, f"G{r}", "", f_in, fill_in)
    r += 1
for col, w in zip("ABCDEFG", [56, 44, 8, 60, 10, 14, 11]):
    st.column_dimensions[col].width = w
st.freeze_panes = "A4"

# =============================================================== Lists
SOIL = [("Kallang Formation (marine clay)", "KF"), ("Fort Canning Boulder Bed", "FCBB"), ("Jurong Formation", "JF"),
        ("Old Alluvium", "OA"), ("Bukit Timah Granite", "BTG"), ("Fill / reclaimed", "FILL"), ("Firm/stiff clay (general)", "FSC")]
lists = {
    "DrainType": [("Precast U-Drain", "PUD"), ("Slabbed-over U-Drain", "SUD"), ("Cast-in-situ Open Drain", "OD"),
                  ("Box Culvert", "BC"), ("Precast Box Culvert", "PBC"), ("Entrance Culvert", "EC"), ("Pipe Culvert", "PC"),
                  ("Common Drain", "CMD"), ("Drop Inlet Chamber", "DIC"), ("Sump", "SP"), ("Silt Trap", "ST"),
                  ("Subsoil Drain", "SSD"), ("Cascade Drain", "CSD")],
    "Category": [("Roadside drain", ""), ("Common drain", ""), ("Outlet / main drain", ""), ("Road crossing culvert", ""),
                 ("Entrance culvert", ""), ("Within development", ""), ("Expressway", "")],
    "Cover": [("Open", ""), ("Covered - RC slab", ""), ("Covered - grating", ""), ("Closed (culvert)", "")],
    "Loading": [("Nil / landscape", "NIL"), ("Footway 5 kPa (SDRE)", "FWY"), ("Drain surcharge 12 kPa (SDRE)", "S12"),
                ("Vehicular LM1 (SS EN 1991-2)", "VEH"), ("Fire engine access (SCDF)", "FEA")],
    "WallType": [("Cantilever RC (toe & heel)", "CANT"), ("L-Wall (heel only)", "LW"), ("Reverse L-Wall (toe only, at boundary)", "RLW"),
                 ("Counterfort RC", "CF"), ("Gravity (mass concrete)", "GRAV"), ("Block / Segmental", "BLK"),
                 ("Gabion", "GAB"), ("Reinforced Earth (MSE)", "MSE")],
    "EmbType": [("Contiguous Bored Pile (CBP)", "CBP"), ("Secant Pile (hard-firm)", "SPW"), ("Secant Pile (hard-soft)", "SPS"),
                ("Steel Sheet Pile", "SSP"), ("Diaphragm Wall", "DW"), ("Soldier Pile & Timber Lagging", "SPL"),
                ("King Post & Lagging", "KPL"), ("Pipe Pile Wall", "PPW")],
    "Support": [("Cantilever (no support)", "CL"), ("1 level strut", "S1"), ("2 levels strut", "S2"), ("3 levels strut", "S3"),
                ("Ground anchors", "GA"), ("Raking strut", "RS"), ("Permanent slabs (top-down)", "TD")],
    "TempPerm": [("Temporary", "T"), ("Permanent", "P")],
    "Soil": SOIL,
    "MHType": [("Standard Manhole (PUB 001C)", "SM"), ("Rectangular Manhole, sewer >= 1000 dia (PUB 004C)", "RM"),
               ("Precast Manhole > 6 m (PUB 006C)", "PM"), ("Drop Manhole", "DM"), ("Inspection Chamber", "IC"),
               ("Valve / Pumping Main Chamber", "VC")],
    "Construct": [("Cast-in-situ", "C"), ("Precast", "P")],
    "Shape": [("Rectangular", ""), ("Circular", "")],
    "ERSSType": [("Steel Sheet Pile Cofferdam", "SSP"), ("Trench Shoring Box", "TSB"), ("Steel Caisson / Liner Plate", "SC"),
                 ("Soldier Pile & Timber Lagging", "SPL"), ("Secant Pile Shaft", "SPS"), ("CBP Shaft", "CBS"),
                 ("Open Cut (battered)", "OC")],
    "Purpose": [("Manhole", ""), ("Jacking pit", ""), ("Receiving pit", ""), ("Sewer trench", ""), ("Sump / chamber", ""), ("Pile cap / footing", "")],
    "LkType": [("Low Covered Linkway", "LCL"), ("High Covered Linkway", "HCL"), ("Modular Covered Linkway (SDRE Rev I)", "MCL")],
    "LkDwg": [("LTA/SDRE25/22/CLW1", ""), ("LTA/SDRE25/22/CLW2", ""), ("LTA/SDRE25/22/CLW3", ""), ("LTA/SDRE25/22/CLW4", ""),
              ("LTA/SDRE25/22/CLW5", ""), ("Non-standard (project design)", "")],
    "LkCols": [("Single column, central", ""), ("Single column, side (cantilever)", ""), ("Double column", "")],
    "Roof": [("Aluminium sheet", ""), ("Metal deck", ""), ("Polycarbonate", ""), ("Aluminium + polycarbonate", "")],
    "Status": [("WIP", ""), ("Checked", ""), ("Approved", ""), ("Superseded", "")],
    "YN": [("Y", ""), ("N", "")],
    "Family": [("Railing", "RLG"), ("Staircase", "STR"), ("POB", "POB"), ("USD", "USD"), ("Bus Shelter", "BS"), ("Others", "OTH")],
    "Unit": [("m", ""), ("nr", ""), ("bay", ""), ("span", ""), ("m2", "")],
    "Grade": [("C25/30", ""), ("C28/35", ""), ("C30/37", ""), ("C32/40", ""), ("C35/45", ""), ("C40/50", "")],
    "Scale": [(10, ""), (20, ""), (25, ""), (50, ""), (100, "")],
}
wl = wb.create_sheet("Lists")
wl["A1"] = "Dropdown lists. Add new items in the blank rows under a list; codes feed the Design IDs."
wl["A1"].font = f_sub
ranges = {}
col = 1
for name, items in lists.items():
    wl.cell(3, col, name).font = f_bold
    wl.cell(3, col + 1, "Code").font = f_bold
    for i, (v, c) in enumerate(items):
        wl.cell(4 + i, col, v).font = f_in
        if c:
            wl.cell(4 + i, col + 1, c).font = f_in
    end = 4 + len(items) + 10
    ranges[name] = (f"Lists!${L(col)}$4:${L(col)}${end}", f"Lists!${L(col)}$4:${L(col+1)}${end}")
    wl.column_dimensions[L(col)].width = 34
    wl.column_dimensions[L(col + 1)].width = 7
    col += 3


def dv_list(ws, name, ref):
    dv = DataValidation(type="list", formula1="=" + ranges[name][0], allow_blank=True)
    dv.error, dv.errorTitle = "Pick a value from the list (add new values on the Lists sheet).", "Not in list"
    ws.add_data_validation(dv)
    dv.add(ref)


def vlk(cell, lst):
    return f'IFERROR(VLOOKUP({cell},{ranges[lst][1]},2,0),"??")'


# =============================================================== register machinery
REGS = []  # (sheet, idx, label, keycol)


def build_register(title, sheetname, note, cols, examples, tname, label):
    ws = wb.create_sheet(sheetname)
    ws["A1"], ws["A2"] = title, note
    ws["A1"].font, ws["A2"].font = f_title, f_sub
    ws["A3"] = ("Yellow = type / pick.  Grey = automatic.  Purple = the design's file set (built from the Design ID).  "
                "Green = tender quantities per unit.  Green-shaded rows = EXAMPLE data, values illustrative: delete or overwrite.")
    ws["A3"].font = f_sub
    idx = {h: L(c) for c, (h, *_r) in enumerate(cols, 1)}
    for c, (h, kind, w, ex) in enumerate(cols, 1):
        cell = ws.cell(HDR, c, h)
        cell.font, cell.alignment, cell.border = f_hdr, wrap, box
        cell.fill = fill_hdr["file" if ex.get("file") else kind]
        ws.column_dimensions[L(c)].width = w
    ws.row_dimensions[HDR].height = 54
    for c, (h, kind, w, ex) in enumerate(cols, 1):
        for r in range(FIRST, LAST + 1):
            cell = ws.cell(r, c)
            cell.border = box
            if "formula" in ex:
                cell.value = ex["formula"](r, idx)
                cell.font, cell.fill = f_calc, fill_calc
            else:
                cell.font, cell.fill = f_in, fill_in
            if "fmt" in ex:
                cell.number_format = ex["fmt"]
        if "dv" in ex:
            dv_list(ws, ex["dv"], f"{L(c)}{FIRST}:{L(c)}{LAST}")
    for i, row in enumerate(examples):
        r = FIRST + i
        for h, v in row.items():
            ws[f"{idx[h]}{r}"] = v
        for c in range(1, len(cols) + 1):
            if "formula" not in cols[c - 1][3]:
                ws.cell(r, c).fill = fill_ex
    ws.freeze_panes = ws[f"C{FIRST}"]
    t = Table(displayName=tname, ref=f"A{HDR}:{L(len(cols))}{LAST}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    ws.sheet_view.zoomScale = 90
    for h in ("SG Compliance Check", "BCA ERSS Submission (indicative)"):
        if h in idx:
            c = idx[h]
            rg = f"{c}{FIRST}:{c}{LAST}"
            ws.conditional_formatting.add(rg, FormulaRule(formula=[f'ISNUMBER(SEARCH("FAIL",{c}{FIRST}))'], fill=fill_red))
            ws.conditional_formatting.add(rg, FormulaRule(formula=[f'ISNUMBER(SEARCH("NOTE",{c}{FIRST}))'], fill=fill_amb))
            ws.conditional_formatting.add(rg, FormulaRule(formula=[f'{c}{FIRST}="OK"'], fill=fill_grn))
    REGS.append((sheetname, idx, label))
    return ws, idx


def C(i, r):
    return lambda h: f"{i[h]}{r}"


def f_comp(keyh, parts_fn):
    def f(r, i):
        cat = "&".join(parts_fn(C(i, r)))
        k = C(i, r)(keyh)
        return f'=IF({k}="","",IF(LEN({cat})=0,"OK",{cat}))'
    return f


def f_sel(keyh):
    def f(r, i):
        c = C(i, r)
        return f'=IF({c(keyh)}="","",IF(AND({c("Status")}="Approved",ISERROR(SEARCH("FAIL",{c("SG Compliance Check")}))),"Y","N"))'
    return f


def f_key(keyh, fields):
    def f(r, i):
        c = C(i, r)
        return f'=IF({c(keyh)}="","",' + '&"|"&'.join(c(h) for h in fields) + f'&"|"&{c("Selectable")})'
    return f


def f_bca(keyh, depth_h, nonrc_cond):
    """Indicative BCA ERSS submission route from excavation / retained depth."""
    def f(r, i):
        c = C(i, r)
        d = c(depth_h)
        thr0 = f'IF({nonrc_cond(c)},{S["bca0n"]},{S["bca0"]})'
        return (f'=IF(OR({c(keyh)}="",{d}=""),"",IF({d}<={thr0},"No ERSS submission",IF({d}<={S["bca1"]},"QP(ST)",'
                f'IF({d}<={S["bca2"]},"QP(ST) + AC","Geotechnical building works: QP(Geo) + AC"))))')
    return f


def grade_chk(c, h, key):
    return f'IF(AND({c(h)}<>"",IFERROR(VALUE(MID({c(h)},2,2)),0)<{S[key]}),"FAIL: grade below C"&{S[key]}&". ","")'


def cover_chk(c):
    return f'IF(AND({c("Cover (mm)")}<>"",{c("Cover (mm)")}<{S["cov"]}),"FAIL: cover < "&{S["cov"]}&" mm. ","")'


def file_cols(folder_key):
    def folder(r, i):
        return f'=IF({i["Design ID"]}{r}="","",{S["rootn"]}&{S[folder_key]}&"\\"&{i["Design ID"]}{r}&"\\")'

    def fn(suffix, cond=None):
        def f(r, i):
            base = f'{i["Design ID"]}{r}&"{suffix}"'
            if cond:
                return f'=IF({i["Design ID"]}{r}="","",IF({i[cond]}{r}="Y",{base},"n/a"))'
            return f'=IF({i["Design ID"]}{r}="","",{base})'
        return f
    return [
        ("Detail Scale (1:n)", "in", 8, {"dv": "Scale"}),
        ("STAAD.Pro Model? (Y/N)", "in", 9, {"dv": "YN"}),
        ("Design Folder", "calc", 46, {"formula": folder, "file": True}),
        ("CAD Detail (.dwg)", "calc", 30, {"formula": fn("_CAD.dwg"), "file": True}),
        ("Drawing (.pdf)", "calc", 30, {"formula": fn("_DWG.pdf"), "file": True}),
        ("STAAD.Pro (.std)", "calc", 30, {"formula": fn("_STAAD.std", "STAAD.Pro Model? (Y/N)"), "file": True}),
        ("Calc Sheet (.xlsx)", "calc", 30, {"formula": fn("_CALC.xlsx"), "file": True}),
        ("Calc Report (.pdf)", "calc", 30, {"formula": fn("_CALC.pdf"), "file": True}),
        ("Files Checked in Folder (Y/N)", "in", 9, {"dv": "YN"}),
    ]


QTY = ["Concrete (m3/unit)", "Rebar (kg/unit)", "Steel (kg/unit)", "Formwork (m2/unit)", "Excavation (m3/unit)"]


def qty_cols():
    return [("Qty Unit", "in", 6, {"dv": "Unit"})] + [(h, "qty", 9, {"fmt": "0.00"}) for h in QTY]


def admin_cols(head):
    return [("Code Basis", "in", 30, {}), ("Rev", "in", 5, {}), ("Status", "in", 10, {"dv": "Status"}),
            ("Drawing No.", "in", 14, {})]


TAIL = [("Limitations / Notes", "in", 36, {}), ("Projects / Tenders Used", "in", 22, {}), ("Designer", "in", 10, {}),
        ("Checker", "in", 10, {}), ("Approved By (QP)", "in", 11, {}), ("Approval Date", "in", 11, {"fmt": "dd-mmm-yy"})]


def ex_common(rev, status, dwg, scale, staad, unit, q, lim, cb, proj=""):
    d = {"Rev": rev, "Status": status, "Drawing No.": dwg, "Detail Scale (1:n)": scale,
         "STAAD.Pro Model? (Y/N)": "Y" if staad else "N", "Files Checked in Folder (Y/N)": "N", "Qty Unit": unit,
         "Limitations / Notes": "EXAMPLE. " + lim, "Projects / Tenders Used": proj, "Code Basis": cb,
         "Designer": "EXAMPLE", "Checker": "EXAMPLE", "Approved By (QP)": "EXAMPLE", "Approval Date": TODAY}
    d.update(dict(zip(QTY, q)))
    return d


# =============================================================== 1. DRAINS
NONHYD = ["Sump", "Drop Inlet Chamber", "Silt Trap", "Pipe Culvert", "Subsoil Drain"]
HYD_S = "Design_Selector!$B$11"
WD, DD = "Internal Width (mm)", "Internal Depth / Clear Height (mm)"


def d_id(r, i):
    c = C(i, r)
    return (f'=IF({c("Drain Type")}="","","DRN-"&{vlk(c("Drain Type"),"DrainType")}&"-"&TEXT({c(WD)},"0000")&"x"&TEXT({c(DD)},"0000")'
            f'&IF(N({c("No. of Cells")})>1,"-"&{c("No. of Cells")}&"C","")&"-"&{vlk(c("Loading"),"Loading")}&"-"&{c("Rev")})')


def d_desc(r, i):
    c = C(i, r)
    return (f'=IF({c("Drain Type")}="","",{c("Drain Type")}&" "&{c(WD)}&"W x "&{c(DD)}&"D"&IF(N({c("No. of Cells")})>1," ("&{c("No. of Cells")}&" cells)","")'
            f'&", "&{c("Category")}&", "&{c("Loading")}&", "&{c("Open / Covered")})')


def d_y(c):
    return f'({c(DD)}/1000*(1-{S["fb"]}))'


def d_q(r, i):
    c = C(i, r)
    W, y = f'({c(WD)}/1000)', d_y(c)
    nonhyd = ",".join(f'{c("Drain Type")}="{t}"' for t in NONHYD)
    return (f'=IF(OR({c("Drain Type")}="",{c(WD)}="",{c(DD)}="",{nonhyd}),"",ROUND(MAX(1,N({c("No. of Cells")}))*(1/{S["nman"]})'
            f'*({W}*{y})*(({W}*{y})/({W}+2*{y}))^(2/3)*(1/{HYD_S})^0.5,3))')


def d_v(r, i):
    c = C(i, r)
    return f'=IF({c("Flow Capacity Q (m3/s)")}="","",ROUND({c("Flow Capacity Q (m3/s)")}/(MAX(1,N({c("No. of Cells")}))*{c(WD)}/1000*{d_y(c)}),2))'


def d_parts(c):
    return [
        f'IF(AND(OR({c("Category")}="Roadside drain",{c("Category")}="Entrance culvert"),OR({c(WD)}<{S["rdW"]},{c(DD)}<{S["rdD"]})),"FAIL: below PUB min "&{S["rdW"]}&"x"&{S["rdD"]}&" for roadside drain. ","")',
        f'IF(AND({c("Category")}="Common drain",{c(WD)}<{S["cdW"]}),"FAIL: common drain W < "&{S["cdW"]}&" mm. ","")',
        grade_chk(c, "Concrete Grade", "fck"), cover_chk(c),
        f'IF(AND({c("Open / Covered")}="Open",{c(DD)}>{S["rail"]}),"NOTE: safety railing required (D > "&{S["rail"]}&" mm). ","")',
        f'IF(AND({c("Category")}="Entrance culvert",{c(DD)}<600),"NOTE: check hinged grating rule for culverts < 600 deep (PUB Add. 3). ","")',
    ]


drain_cols = [
    ("Design ID", "calc", 32, {"formula": d_id}), ("Description", "calc", 44, {"formula": d_desc}),
    ("Drain Type", "in", 18, {"dv": "DrainType"}), ("Category", "in", 17, {"dv": "Category"}),
    ("Open / Covered", "in", 14, {"dv": "Cover"}), (WD, "in", 9, {"fmt": "0"}), (DD, "in", 10, {"fmt": "0"}),
    ("No. of Cells", "in", 6, {"fmt": "0"}), ("Wall Thk (mm)", "in", 7, {"fmt": "0"}), ("Base Thk (mm)", "in", 7, {"fmt": "0"}),
    ("Top Slab Thk (mm)", "in", 7, {"fmt": "0"}), ("Loading", "in", 20, {"dv": "Loading"}),
    ("Fill over Top / Surcharge", "in", 14, {}), ("Concrete Grade", "in", 8, {"dv": "Grade"}), ("Cover (mm)", "in", 7, {"fmt": "0"}),
    ("Wall Main Bar", "in", 10, {}), ("Base Main Bar", "in", 10, {}), ("Top Slab Bar", "in", 10, {}), ("Distribution Bar", "in", 10, {}),
    *admin_cols("Drain"), *file_cols("f_dr"), *qty_cols(),
    ("Flow Capacity Q (m3/s)", "calc", 10, {"formula": d_q, "fmt": "0.000"}),
    ("Velocity at Capacity (m/s)", "calc", 9, {"formula": d_v, "fmt": "0.00"}),
    ("SG Compliance Check", "calc", 44, {"formula": f_comp("Drain Type", d_parts)}),
    *TAIL, ("Selectable", "calc", 8, {"formula": f_sel("Drain Type")}),
    ("Selector Key", "calc", 34, {"formula": f_key("Drain Type", ["Drain Type", WD, DD, "Loading"])}),
]
S12, FWY, VEH = "Drain surcharge 12 kPa (SDRE)", "Footway 5 kPa (SDRE)", "Vehicular LM1 (SS EN 1991-2)"
CB_DR = "PUB COP SWD 7th Ed + Add.3; LTA SDRE Rev I Ch.2; SS EN 1992-1-1"
CB_CUL = "PUB COP SWD; LTA CDC C&S A3; SS EN 1991-2 LM1; SS EN 1992-1-1"


def dr(t, cat, cov, w, d, cells, wt, bt, ts, load, fill, grade, cvr, wbar, bbar, tbar, dbar, rev, stt, dwg, scale, staad, unit, q, lim, cb=CB_DR, proj=""):
    e = {"Drain Type": t, "Category": cat, "Open / Covered": cov, WD: w, DD: d, "No. of Cells": cells, "Wall Thk (mm)": wt,
         "Base Thk (mm)": bt, "Top Slab Thk (mm)": ts, "Loading": load, "Fill over Top / Surcharge": fill, "Concrete Grade": grade,
         "Cover (mm)": cvr, "Wall Main Bar": wbar, "Base Main Bar": bbar, "Top Slab Bar": tbar, "Distribution Bar": dbar}
    e.update(ex_common(rev, stt, dwg, scale, staad, unit, q, lim, cb, proj))
    return e


drain_ex = [
    dr("Precast U-Drain", "Roadside drain", "Open", 600, 600, 1, 100, 100, None, S12, "12 kPa", "C32/40", 40, "B10 mesh", "B10 mesh", "-", "-", "A", "Approved", "STD-DR-001", 20, False, "m", (.18, 14, 0, 0, 1.3), "Per LTA SDRE precast U-drain; confirm supplier unit."),
    dr("Precast U-Drain", "Roadside drain", "Open", 900, 900, 1, 125, 125, None, S12, "12 kPa", "C32/40", 40, "B12 mesh", "B12 mesh", "-", "-", "A", "Approved", "STD-DR-002", 20, False, "m", (.34, 24, 0, 0, 2.2), ""),
    dr("Slabbed-over U-Drain", "Roadside drain", "Covered - RC slab", 600, 750, 1, 125, 125, 125, FWY, "Footway 5 kPa", "C32/40", 40, "T10@150", "T10@150", "T10@150", "T10@200", "A", "Approved", "STD-DR-003", 20, False, "m", (.43, 38, 0, 2.6, 1.8), "Under footway only."),
    dr("Cast-in-situ Open Drain", "Outlet / main drain", "Open", 1200, 1200, 1, 200, 200, None, S12, "12 kPa", "C32/40", 50, "T13@150", "T13@150", "-", "T10@200", "A", "Approved", "STD-DR-004", 25, False, "m", (.72, 72, 0, 4.8, 3.6), "D > 1.0 m: add railing STD-RLG-001."),
    dr("Cast-in-situ Open Drain", "Outlet / main drain", "Open", 2000, 1800, 1, 250, 300, None, S12, "12 kPa", "C32/40", 50, "T16@150", "T16@150", "-", "T13@200", "A", "Checked", "STD-DR-005", 25, True, "m", (1.65, 150, 0, 7.2, 7.0), "Awaiting QP approval."),
    dr("Box Culvert", "Road crossing culvert", "Closed (culvert)", 1500, 1200, 1, 250, 250, 250, VEH, "0.5 - 2.0 m fill", "C32/40", 50, "T16@150", "T16@150", "T16@150", "T13@200", "A", "Approved", "STD-BC-001", 25, True, "m", (1.88, 210, 0, 7.4, 6.5), "STAAD frame: LM1 + earth pressure.", CB_CUL, "EXAMPLE Tender T-2026-01"),
    dr("Box Culvert", "Road crossing culvert", "Closed (culvert)", 2400, 1800, 2, 300, 300, 300, VEH, "0.5 - 2.0 m fill", "C32/40", 50, "T20@150", "T20@150", "T20@150", "T13@200", "A", "Approved", "STD-BC-002", 50, True, "m", (5.2, 610, 0, 20.5, 16.0), "Twin cell.", CB_CUL),
    dr("Entrance Culvert", "Entrance culvert", "Covered - grating", 600, 600, 1, 150, 150, None, VEH, "Heavy duty grating", "C32/40", 40, "T13@150", "T13@150", "-", "T10@200", "A", "Approved", "STD-EC-001", 20, False, "m", (.41, 45, 0, 2.9, 1.9), "Hinged heavy-duty grating."),
    dr("Common Drain", "Common drain", "Covered - grating", 450, 450, 1, 100, 100, None, FWY, "-", "C32/40", 40, "T10@200", "T10@200", "-", "T10@250", "A", "Approved", "STD-CMD-001", 20, False, "m", (.15, 12, 0, 1.9, .9), "Within development only."),
    dr("Precast U-Drain", "Roadside drain", "Open", 450, 450, 1, 100, 100, None, S12, "12 kPa", "C25/30", 40, "B10 mesh", "B10 mesh", "-", "-", "A", "WIP", "STD-DR-099", 20, False, "m", (.14, 11, 0, 0, 1.0), "Shows a design the SG check rejects."),
    dr("Drop Inlet Chamber", "Roadside drain", "Covered - grating", 600, 900, 1, 150, 150, None, S12, "-", "C32/40", 40, "T10@150", "T10@150", "-", "-", "A", "Approved", "STD-DIC-001", 20, False, "nr", (.55, 42, 0, 4.2, 1.6), "250 dia uPVC outlet."),
    dr("Sump", "Roadside drain", "Covered - grating", 900, 1200, 1, 200, 200, None, S12, "-", "C32/40", 40, "T13@150", "T13@150", "-", "T10@200", "A", "Approved", "STD-SP-001", 20, False, "nr", (1.25, 110, 0, 9.8, 3.2), ""),
]
wsd, di = build_register("Drainage Standard Design Register (Singapore)", "Drain_Register",
                         "Drains, culverts, sumps, chambers. Compliance uses SG_Criteria (PUB COP SWD, LTA SDRE). Quantities per Qty Unit.",
                         drain_cols, drain_ex, "tblDrain", "Drain")

# =============================================================== 2. RC / GRAVITY RETAINING WALLS
RWH = "Retained Height (m)"
NONRC_RW = ("Block / Segmental", "Gabion", "Reinforced Earth (MSE)")


def rw_id(r, i):
    c = C(i, r)
    return (f'=IF({c("Wall Type")}="","","RW-"&{vlk(c("Wall Type"),"WallType")}&"-H"&TEXT({c(RWH)}*1000,"0000")&"-S"&{c("Surcharge (kPa)")}'
            f'&IF({c("GWT Behind Wall (Y/N)")}="Y","-GW","")&"-"&{c("Rev")})')


def rw_desc(r, i):
    c = C(i, r)
    return (f'=IF({c("Wall Type")}="","",{c("Wall Type")}&", H = "&TEXT({c(RWH)},"0.0")&" m, q = "&{c("Surcharge (kPa)")}&" kPa"'
            f'&IF({c("GWT Behind Wall (Y/N)")}="Y",", with GWT","")&", "&{c("Founding Soil")})')


def rw_parts(c):
    return [grade_chk(c, "Concrete Grade", "fck"), cover_chk(c),
            f'IF(AND({c("FOS Sliding")}<>"",{c("FOS Sliding")}<1),"FAIL: sliding. ","")',
            f'IF(AND({c("Wall Type")}="Reverse L-Wall (toe only, at boundary)",N({c("Heel (mm)")})>0),"NOTE: reverse L-wall should have no heel. ","")',
            f'IF(AND({c("Wall Type")}="L-Wall (heel only)",N({c("Toe (mm)")})>0),"NOTE: L-wall should have no toe. ","")',
            f'IF(AND({c("Surcharge (kPa)")}<>"",{c("Surcharge (kPa)")}<{S["qrw"]}),"NOTE: q below LTA min "&{S["qrw"]}&" kPa (not for LTA works). ","")']


rw_cols = [
    ("Design ID", "calc", 28, {"formula": rw_id}), ("Description", "calc", 46, {"formula": rw_desc}),
    ("Wall Type", "in", 22, {"dv": "WallType"}), (RWH, "in", 8, {"fmt": "0.00"}), ("Surcharge (kPa)", "in", 8, {"fmt": "0"}),
    ("GWT Behind Wall (Y/N)", "in", 8, {"dv": "YN"}), ("Backfill Slope (deg)", "in", 7, {"fmt": "0"}),
    ("Founding Soil", "in", 20, {"dv": "Soil"}), ("phi' (deg)", "in", 6, {"fmt": "0"}), ("gamma (kN/m3)", "in", 7, {"fmt": "0.0"}),
    ("Stem Top Thk (mm)", "in", 7, {"fmt": "0"}), ("Stem Bottom Thk (mm)", "in", 7, {"fmt": "0"}),
    ("Base Width (mm)", "in", 7, {"fmt": "0"}), ("Toe (mm)", "in", 6, {"fmt": "0"}), ("Heel (mm)", "in", 6, {"fmt": "0"}),
    ("Base Thk (mm)", "in", 6, {"fmt": "0"}), ("Shear Key (mm)", "in", 7, {}), ("Concrete Grade", "in", 8, {"dv": "Grade"}),
    ("Cover (mm)", "in", 6, {"fmt": "0"}), ("Stem Main Bar", "in", 10, {}), ("Base Main Bar", "in", 10, {}),
    ("Drainage (weepholes / filter)", "in", 16, {}), ("FOS Sliding", "in", 6, {"fmt": "0.00"}),
    ("FOS Overturning", "in", 7, {"fmt": "0.00"}), ("Max Bearing (kPa)", "in", 7, {"fmt": "0"}),
    *admin_cols("RW"), *file_cols("f_rw"), *qty_cols(),
    ("BCA ERSS Submission (indicative)", "calc", 22, {"formula": f_bca("Wall Type", RWH, lambda c: "OR(" + ",".join(f'{c("Wall Type")}="{t}"' for t in NONRC_RW) + ")")}),
    ("SG Compliance Check", "calc", 40, {"formula": f_comp("Wall Type", rw_parts)}),
    *TAIL, ("Selectable", "calc", 8, {"formula": f_sel("Wall Type")}),
    ("Selector Key", "calc", 34, {"formula": f_key("Wall Type", ["Wall Type", RWH, "Surcharge (kPa)", "GWT Behind Wall (Y/N)"])}),
]
CB_RW = "SS EN 1997-1 DA1 + NA; SS EN 1992-1-1; LTA CDC C&S A3"
CANT, LW, RLW = "Cantilever RC (toe & heel)", "L-Wall (heel only)", "Reverse L-Wall (toe only, at boundary)"
FSC = "Firm/stiff clay (general)"


def rw(t, h, q, gw, soil, phi, stt_, stb, B, toe, heel, bt, sbar, bbar, fs, fo, qb, rev, status, dwg, scale, staad, qq, lim):
    e = {"Wall Type": t, RWH: h, "Surcharge (kPa)": q, "GWT Behind Wall (Y/N)": gw, "Backfill Slope (deg)": 0,
         "Founding Soil": soil, "phi' (deg)": phi, "gamma (kN/m3)": 19.0, "Stem Top Thk (mm)": stt_, "Stem Bottom Thk (mm)": stb,
         "Base Width (mm)": B, "Toe (mm)": toe, "Heel (mm)": heel, "Base Thk (mm)": bt, "Shear Key (mm)": "-",
         "Concrete Grade": "C32/40", "Cover (mm)": 50, "Stem Main Bar": sbar, "Base Main Bar": bbar,
         "Drainage (weepholes / filter)": "75 dia @ 1.5 m c/c + geotextile", "FOS Sliding": fs, "FOS Overturning": fo, "Max Bearing (kPa)": qb}
    e.update(ex_common(rev, status, dwg, scale, staad, "m", qq, lim, CB_RW))
    return e


rw_ex = [
    rw(LW, 1.0, 25, "N", FSC, 26, 200, 200, 1300, 0, 1100, 250, "T10@200", "T10@200", 1.6, 2.4, 65, "A", "Approved", "STD-RW-001", 25, False, (.53, 38, 0, 2.5, 1.9), ""),
    rw(LW, 2.0, 25, "N", FSC, 26, 200, 300, 2300, 0, 2000, 350, "T13@150", "T13@150", 1.5, 2.1, 100, "A", "Approved", "STD-RW-011", 25, False, (1.35, 98, 0, 4.7, 4.0), ""),
    rw(RLW, 1.5, 25, "N", FSC, 26, 200, 250, 1700, 1450, 0, 300, "T13@150", "T13@150", 1.5, 1.8, 120, "A", "Approved", "STD-RW-012", 25, False, (.92, 70, 0, 3.6, 2.9), "For walls on the site boundary (no heel)."),
    rw(CANT, 1.5, 25, "N", FSC, 26, 200, 250, 1650, 400, 1000, 300, "T13@200", "T13@200", 1.6, 2.3, 80, "A", "Approved", "STD-RW-002", 25, False, (.9, 66, 0, 3.6, 2.8), ""),
    rw(CANT, 2.0, 25, "N", FSC, 26, 200, 300, 2100, 500, 1300, 350, "T13@150", "T13@150", 1.5, 2.2, 95, "A", "Approved", "STD-RW-003", 25, False, (1.28, 94, 0, 4.7, 3.7), ""),
    rw(CANT, 2.5, 25, "N", FSC, 26, 200, 350, 2500, 600, 1550, 400, "T16@150", "T16@150", 1.5, 2.1, 110, "A", "Approved", "STD-RW-004", 25, False, (1.7, 128, 0, 5.8, 4.6), ""),
    rw(CANT, 3.0, 25, "N", FSC, 26, 200, 400, 2950, 700, 1850, 450, "T16@125", "T16@150", 1.5, 2.0, 128, "A", "Approved", "STD-RW-005", 25, False, (2.2, 170, 0, 6.9, 5.7), ""),
    rw(CANT, 3.5, 25, "N", FSC, 26, 250, 450, 3350, 800, 2100, 500, "T20@150", "T16@125", 1.5, 2.0, 142, "A", "Approved", "STD-RW-006", 50, False, (2.78, 222, 0, 8.0, 6.9), ""),
    rw(CANT, 3.0, 25, "Y", FSC, 26, 250, 450, 3300, 800, 2050, 500, "T20@150", "T20@150", 1.5, 2.0, 145, "A", "Approved", "STD-RW-007", 25, False, (2.58, 210, 0, 7.1, 6.4), "GWT at 1/2 H behind wall."),
    rw(CANT, 4.5, 25, "N", FSC, 26, 250, 550, 4100, 1000, 2550, 600, "T25@150", "T20@125", 1.5, 2.0, 170, "A", "Checked", "STD-RW-008", 50, True, (4.05, 345, 0, 10.2, 9.3), "Awaiting QP approval."),
    rw("Counterfort RC", 6.0, 25, "N", "Old Alluvium", 30, 300, 300, 4500, 1000, 3200, 600, "T20@150", "T20@150", 1.5, 2.0, 210, "A", "WIP", "STD-RW-013", 50, True, (5.9, 520, 0, 16.0, 14.5), "Counterforts at 3.0 m c/c."),
    rw("Gravity (mass concrete)", 1.5, 10, "N", "Old Alluvium", 32, 400, 700, 1100, 0, 0, 0, "-", "-", 1.6, 2.5, 80, "A", "Approved", "STD-RW-009", 25, False, (.95, 0, 0, 3.2, 1.9), "Mass concrete, no rebar."),
]
wsr, ri = build_register("RC & Gravity Retaining Wall Register (Singapore)", "RW_Register",
                         "Cantilever (toe & heel), L-wall (heel only), reverse L-wall (toe only), counterfort, gravity, block, gabion, MSE. Embedded walls (CBP, secant, sheet pile, D-wall) go in Embedded_Wall_Register.",
                         rw_cols, rw_ex, "tblRW", "Ret. Wall")

# =============================================================== 3. EMBEDDED WALLS (permanent or temporary ERSS, linear)
EXD = "Excavation Depth (m)"


def ew_id(r, i):
    c = C(i, r)
    return (f'=IF({c("Wall Type")}="","","EW-"&{vlk(c("Wall Type"),"EmbType")}&"-D"&TEXT({c(EXD)}*1000,"0000")&"-"&{vlk(c("Support System"),"Support")}'
            f'&"-"&{vlk(c("Soil Profile"),"Soil")}&"-"&{vlk(c("Temporary / Permanent"),"TempPerm")}&"-"&{c("Rev")})')


def ew_desc(r, i):
    c = C(i, r)
    return (f'=IF({c("Wall Type")}="","",{c("Temporary / Permanent")}&" "&{c("Wall Type")}&", excavation "&TEXT({c(EXD)},"0.0")&" m, "'
            f'&{c("Support System")}&", "&{c("Soil Profile")}&", "&{c("Pile Dia / Section")})')


def ew_parts(c):
    return [grade_chk(c, "Concrete Grade", "fck"),
            f'IF(AND(LEFT({c("Wall Type")},3)="Con",N({c("GWT (m below ground)")})<{c(EXD)},{c("Soil Profile")}="Kallang Formation (marine clay)"),"NOTE: CBP is not watertight below GWT in marine clay; consider secant / grouting. ","")',
            f'IF(AND({c("Wall Type")}="Steel Sheet Pile",{c("Temporary / Permanent")}="Permanent"),"NOTE: permanent sheet pile needs corrosion allowance. ","")',
            f'IF(AND({c("Support System")}="Cantilever (no support)",{c(EXD)}>6),"NOTE: cantilever embedded wall > 6 m, check deflection. ","")',
            f'IF(AND({c("Max Wall Deflection (mm)")}<>"",{c("Allowable Deflection (mm)")}<>"",{c("Max Wall Deflection (mm)")}>{c("Allowable Deflection (mm)")}),"FAIL: deflection exceeds allowable. ","")']


ew_cols = [
    ("Design ID", "calc", 34, {"formula": ew_id}), ("Description", "calc", 50, {"formula": ew_desc}),
    ("Wall Type", "in", 22, {"dv": "EmbType"}), ("Temporary / Permanent", "in", 11, {"dv": "TempPerm"}),
    (EXD, "in", 8, {"fmt": "0.0"}), ("Support System", "in", 16, {"dv": "Support"}), ("Soil Profile", "in", 22, {"dv": "Soil"}),
    ("GWT (m below ground)", "in", 8, {"fmt": "0.0"}), ("Surcharge (kPa)", "in", 8, {"fmt": "0"}),
    ("Pile Dia / Section", "in", 14, {}), ("Spacing / Overlap (mm)", "in", 10, {}), ("Wall Thk (mm)", "in", 7, {}),
    ("Toe Embedment (m)", "in", 8, {"fmt": "0.0"}), ("Pile / Wall Length (m)", "in", 8, {"fmt": "0.0"}),
    ("Main Bar / Cage", "in", 14, {}), ("Capping Beam (W x D mm)", "in", 12, {}), ("Waler Section", "in", 12, {}),
    ("Strut Section & Spacing", "in", 16, {}), ("Max Wall Deflection (mm)", "in", 8, {"fmt": "0"}),
    ("Allowable Deflection (mm)", "in", 8, {"fmt": "0"}), ("Concrete Grade", "in", 8, {"dv": "Grade"}),
    ("Instrumentation", "in", 18, {}), *admin_cols("EW"), *file_cols("f_ew"), *qty_cols(),
    ("BCA ERSS Submission (indicative)", "calc", 24, {"formula": f_bca("Wall Type", EXD, lambda c: f'{c("Wall Type")}="Soldier Pile & Timber Lagging"')}),
    ("SG Compliance Check", "calc", 42, {"formula": f_comp("Wall Type", ew_parts)}),
    *TAIL, ("Selectable", "calc", 8, {"formula": f_sel("Wall Type")}),
    ("Selector Key", "calc", 40, {"formula": f_key("Wall Type", ["Wall Type", EXD, "Support System", "Soil Profile", "Temporary / Permanent"])}),
]
CB_EW = "SS EN 1997-1 DA1 + NA; SS EN 1992-1-1; BCA ERSS requirements"
KF, OA, JF = "Kallang Formation (marine clay)", "Old Alluvium", "Jurong Formation"


def ew(t, tp, d, sup, soil, gwt, q, sec, spc, th, toe, ln, bar, cap, wal, strut, dfl, adf, inst, rev, status, dwg, scale, staad, qq, lim):
    e = {"Wall Type": t, "Temporary / Permanent": tp, EXD: d, "Support System": sup, "Soil Profile": soil,
         "GWT (m below ground)": gwt, "Surcharge (kPa)": q, "Pile Dia / Section": sec, "Spacing / Overlap (mm)": spc,
         "Wall Thk (mm)": th, "Toe Embedment (m)": toe, "Pile / Wall Length (m)": ln, "Main Bar / Cage": bar,
         "Capping Beam (W x D mm)": cap, "Waler Section": wal, "Strut Section & Spacing": strut,
         "Max Wall Deflection (mm)": dfl, "Allowable Deflection (mm)": adf, "Concrete Grade": "C32/40" if t != "Steel Sheet Pile" else "",
         "Instrumentation": inst}
    e.update(ex_common(rev, status, dwg, scale, staad, "m", qq, lim, CB_EW))
    return e


ew_ex = [
    ew("Contiguous Bored Pile (CBP)", "Permanent", 3.0, "Cantilever (no support)", OA, 2.0, 20, "600 dia", "750 c/c", "-", 6.0, 9.0, "12T25, R10@200", "900 x 800", "-", "-", 22, 35, "Inclinometer @ 20 m", "A", "Approved", "STD-EW-001", 50, True, (2.6, 260, 0, 1.6, 0.6), "Plus lagging / shotcrete between piles."),
    ew("Contiguous Bored Pile (CBP)", "Permanent", 5.0, "1 level strut", OA, 2.0, 20, "800 dia", "950 c/c", "-", 7.0, 12.0, "16T32, R12@150", "1100 x 900", "UB 610x229", "CHS 508x12.5 @ 6 m", 28, 40, "Inclinometer, strain gauges", "A", "Approved", "STD-EW-002", 50, True, (6.4, 720, 180, 1.8, 0.9), ""),
    ew("Secant Pile (hard-firm)", "Permanent", 6.0, "1 level strut", KF, 1.0, 20, "1000 dia", "200 overlap", "-", 12.0, 18.0, "Primary 20T32", "1300 x 1000", "UB 762x267", "CHS 610x16 @ 6 m", 30, 40, "Inclinometer, piezometer, settlement markers", "A", "Approved", "STD-EW-003", 50, True, (14.1, 1150, 220, 2.2, 1.2), "Watertight; toe into Old Alluvium."),
    ew("Steel Sheet Pile", "Temporary", 4.0, "1 level strut", KF, 1.0, 20, "FSP IV", "-", "-", 8.0, 12.0, "-", "-", "UB 457x191", "CHS 406x12.5 @ 5 m", 35, 50, "Inclinometer, piezometer", "A", "Approved", "STD-EW-004", 50, True, (0, 0, 2280, 0, 0), "Per BCA sheet pile integrity guideline (Nov 2025)."),
    ew("Diaphragm Wall", "Permanent", 15.0, "3 levels strut", KF, 1.0, 20, "-", "Panels 6 m", 1000, 20.0, 35.0, "T40@150 both faces", "1200 x 1200", "Twin UB 914", "Twin H400 @ 8 m", 45, 60, "Full instrumentation plan", "A", "WIP", "STD-EW-005", 100, True, (35, 4200, 900, 0, 0), "Indicative only; project-specific design."),
]
wse, ei = build_register("Embedded Retaining Wall / ERSS Register (Singapore)", "Embedded_Wall_Register",
                         "CBP, secant, sheet pile, diaphragm wall, soldier pile. Tabulated by excavation depth, support system and soil profile. Standard designs are tender-stage typicals for a soil profile class; site-specific design always required.",
                         ew_cols, ew_ex, "tblEmbedded", "Embedded Wall")

# =============================================================== 4. SEWER MANHOLES (permanent)
MHL, MHW, DMIN, DMAX, PIPE = "Internal Length / Dia (mm)", "Internal Width (mm)", "Min Depth (m)", "Max Depth (m)", "Max Incoming Pipe Dia (mm)"


def mh_id(r, i):
    c = C(i, r)
    return (f'=IF({c("Manhole Type")}="","","MH-"&{vlk(c("Manhole Type"),"MHType")}&"-"&TEXT({c(MHL)},"0000")&IF({c("Shape")}="Circular","D","x"&TEXT({c(MHW)},"0000"))'
            f'&"-D"&TEXT({c(DMAX)}*1000,"0000")&"-"&{vlk(c("Construction"),"Construct")}&"-"&{c("Rev")})')


def mh_desc(r, i):
    c = C(i, r)
    return (f'=IF({c("Manhole Type")}="","",{c("Manhole Type")}&", "&{c("Construction")}&", "&IF({c("Shape")}="Circular",{c(MHL)}&" dia",{c(MHL)}&" x "&{c(MHW)})'
            f'&", depth "&TEXT({c(DMIN)},"0.0")&"-"&TEXT({c(DMAX)},"0.0")&" m, pipe up to "&{c(PIPE)}&" dia")')


def mh_parts(c):
    t = c("Manhole Type")
    return [grade_chk(c, "Concrete Grade", "fck"), cover_chk(c),
            f'IF(AND({c(DMIN)}<>"",{c(DMIN)}<{S["mhd"]}),"NOTE: min depth below PUB "&{S["mhd"]}&" m. ","")',
            f'IF(AND(LEFT({t},8)="Standard",{c(DMAX)}>{S["mh6"]}),"FAIL: PUB 001C covers manholes up to "&{S["mh6"]}&" m deep only. ","")',
            f'IF(AND(LEFT({t},11)="Rectangular",{c(PIPE)}<{S["mhbig"]}),"NOTE: 004C is for sewers >= "&{S["mhbig"]}&" dia. ","")',
            f'IF(AND(LEFT({t},7)="Precast",{c(DMAX)}<={S["mh6"]}),"NOTE: 006C is for manholes > "&{S["mh6"]}&" m deep. ","")',
            f'IF(AND({c(PIPE)}<>"",{c(PIPE)}<{S["sdia"]},{t}<>"Inspection Chamber"),"NOTE: public sewer min dia "&{S["sdia"]}&" mm. ","")',
            f'IF(AND({c("Uplift FOS")}<>"",{c("Uplift FOS")}<{S["upl"]}),"FAIL: uplift FOS below "&{S["upl"]}&". ","")']


mh_cols = [
    ("Design ID", "calc", 32, {"formula": mh_id}), ("Description", "calc", 54, {"formula": mh_desc}),
    ("Manhole Type", "in", 26, {"dv": "MHType"}), ("Construction", "in", 11, {"dv": "Construct"}), ("Shape", "in", 10, {"dv": "Shape"}),
    (MHL, "in", 9, {"fmt": "0"}), (MHW, "in", 8, {"fmt": "0"}), (DMIN, "in", 7, {"fmt": "0.0"}), (DMAX, "in", 7, {"fmt": "0.0"}),
    (PIPE, "in", 8, {"fmt": "0"}), ("Wall Thk (mm)", "in", 7, {"fmt": "0"}), ("Base Thk (mm)", "in", 7, {"fmt": "0"}),
    ("Roof Slab Thk (mm)", "in", 7, {"fmt": "0"}), ("Cover & Frame", "in", 16, {}), ("Access (step irons / ladder)", "in", 14, {}),
    ("Benching / Channel", "in", 14, {}), ("Loading", "in", 20, {"dv": "Loading"}), ("Uplift FOS", "in", 7, {"fmt": "0.00"}),
    ("Concrete Grade", "in", 8, {"dv": "Grade"}), ("Cover (mm)", "in", 6, {"fmt": "0"}), ("Wall Bar", "in", 10, {}),
    ("Base Bar", "in", 10, {}), ("Roof Slab Bar", "in", 10, {}), ("PUB Std Drawing Ref", "in", 16, {}),
    ("Typical Temporary ERSS (Design ID)", "in", 30, {}), *admin_cols("MH"), *file_cols("f_mh"), *qty_cols(),
    ("SG Compliance Check", "calc", 44, {"formula": f_comp("Manhole Type", mh_parts)}),
    *TAIL, ("Selectable", "calc", 8, {"formula": f_sel("Manhole Type")}),
    ("Selector Key", "calc", 40, {"formula": f_key("Manhole Type", ["Manhole Type", DMAX, MHL])}),
]
CB_MH = "PUB COP SSW 3rd Ed (Mar 2025); PUB/WRN/STD drawings; SS EN 1992-1-1"
SMH, RMH, PMH = "Standard Manhole (PUB 001C)", "Rectangular Manhole, sewer >= 1000 dia (PUB 004C)", "Precast Manhole > 6 m (PUB 006C)"


def mh(t, cons, shp, l, w, dmin, dmax, pipe, wt, bt, rt, upl, wbar, bbar, rbar, ref, erss, rev, status, dwg, scale, staad, qq, lim):
    e = {"Manhole Type": t, "Construction": cons, "Shape": shp, MHL: l, MHW: w, DMIN: dmin, DMAX: dmax, PIPE: pipe,
         "Wall Thk (mm)": wt, "Base Thk (mm)": bt, "Roof Slab Thk (mm)": rt, "Cover & Frame": "Heavy duty DI, per PUB std",
         "Access (step irons / ladder)": "Step irons @ 300", "Benching / Channel": "Per PUB std", "Loading": VEH, "Uplift FOS": upl,
         "Concrete Grade": "C32/40", "Cover (mm)": 40, "Wall Bar": wbar, "Base Bar": bbar, "Roof Slab Bar": rbar,
         "PUB Std Drawing Ref": ref, "Typical Temporary ERSS (Design ID)": erss}
    e.update(ex_common(rev, status, dwg, scale, staad, "nr", qq, lim, CB_MH))
    return e


mh_ex = [
    mh(SMH, "Cast-in-situ", "Rectangular", 1200, 900, 1.5, 3.0, 300, 200, 250, 200, 1.3, "T13@200", "T13@200", "T13@150", "PUB/WRN/STD/001C", "ERSS-TSB-3000x2500-D3000-FSC-A", "A", "Approved", "STD-MH-001", 25, False, (3.4, 260, 0, 26, 28), "Quantities at max depth."),
    mh(SMH, "Cast-in-situ", "Rectangular", 1500, 1200, 3.0, 6.0, 600, 250, 300, 250, 1.2, "T16@200", "T16@200", "T16@150", "PUB/WRN/STD/001C", "ERSS-SSP-4000x3500-D6000-KF-A", "A", "Approved", "STD-MH-002", 25, True, (9.8, 820, 0, 62, 92), "Quantities at max depth."),
    mh(RMH, "Cast-in-situ", "Rectangular", 2400, 1800, 3.0, 6.0, 1200, 300, 400, 300, 1.15, "T20@150", "T20@150", "T20@150", "PUB/WRN/STD/004C", "ERSS-SSP-4000x3500-D6000-KF-A", "A", "Approved", "STD-MH-003", 50, True, (19.5, 1900, 0, 95, 140), ""),
    mh(PMH, "Precast", "Circular", 1800, None, 6.0, 10.0, 600, 200, 400, 250, 1.25, "Precast rings", "T20@150", "T16@150", "PUB/WRN/STD/006C + 003B", "ERSS-SC-3000x3000-D10000-OA-A", "A", "Checked", "STD-MH-004", 50, True, (12.0, 950, 0, 20, 85), "Awaiting QP approval."),
    mh(SMH, "Cast-in-situ", "Rectangular", 1200, 900, 1.5, 7.5, 300, 200, 250, 200, 1.3, "T13@200", "T13@200", "T13@150", "PUB/WRN/STD/001C", "", "A", "WIP", "STD-MH-099", 25, False, (5.0, 380, 0, 40, 60), "Shows the 001C depth check."),
]
wsm, mi = build_register("Sewer Manhole Register (Singapore, permanent)", "Manhole_Register",
                         "Permanent manholes per PUB COP SSW 3rd Ed and PUB/WRN/STD drawings. One row per size and depth range. Link the typical temporary ERSS pit design.",
                         mh_cols, mh_ex, "tblManhole", "Manhole")

# =============================================================== 5. TEMPORARY ERSS PITS (per nr)
PL, PW = "Pit Plan Length (mm)", "Pit Plan Width (mm)"
NONRC_EP = ("Soldier Pile & Timber Lagging", "Trench Shoring Box", "Open Cut (battered)")


def ep_id(r, i):
    c = C(i, r)
    return (f'=IF({c("ERSS Type")}="","","ERSS-"&{vlk(c("ERSS Type"),"ERSSType")}&"-"&TEXT({c(PL)},"0000")&"x"&TEXT({c(PW)},"0000")'
            f'&"-D"&TEXT({c(EXD)}*1000,"0000")&"-"&{vlk(c("Soil Profile"),"Soil")}&"-"&{c("Rev")})')


def ep_desc(r, i):
    c = C(i, r)
    return (f'=IF({c("ERSS Type")}="","",{c("ERSS Type")}&" for "&{c("Purpose")}&", "&{c(PL)}&" x "&{c(PW)}&" pit, "&TEXT({c(EXD)},"0.0")&" m deep, "'
            f'&{c("Soil Profile")})')


def ep_parts(c):
    return [f'IF(AND({c("Base Stability FOS")}<>"",{c("Base Stability FOS")}<{S["fosb"]}),"FAIL: base stability FOS < "&{S["fosb"]}&". ","")',
            f'IF(AND({c("Purpose")}="Manhole",{c(EXD)}>{S["bca2"]}),"NOTE: BCA lists sewer manholes (pipe <= 2 m dia) > 6 m as non-GBW: QP(ST) + AC. ","")',
            f'IF(AND({c("ERSS Type")}="Trench Shoring Box",{c("Soil Profile")}="Kallang Formation (marine clay)"),"NOTE: check base heave / ground loss for shoring box in marine clay. ","")',
            f'IF(AND({c("ERSS Type")}="Steel Sheet Pile Cofferdam",N({c("Toe Embedment (m)")})<{c(EXD)}*0.5),"NOTE: short toe embedment, check. ","")']


ep_cols = [
    ("Design ID", "calc", 36, {"formula": ep_id}), ("Description", "calc", 54, {"formula": ep_desc}),
    ("ERSS Type", "in", 22, {"dv": "ERSSType"}), ("Purpose", "in", 14, {"dv": "Purpose"}),
    (PL, "in", 8, {"fmt": "0"}), (PW, "in", 8, {"fmt": "0"}), (EXD, "in", 8, {"fmt": "0.0"}), ("Soil Profile", "in", 22, {"dv": "Soil"}),
    ("GWT (m below ground)", "in", 8, {"fmt": "0.0"}), ("Surcharge (kPa)", "in", 8, {"fmt": "0"}),
    ("Sheet Pile / Wall Section", "in", 14, {}), ("Pile Length (m)", "in", 7, {"fmt": "0.0"}), ("Toe Embedment (m)", "in", 7, {"fmt": "0.0"}),
    ("Strut Levels", "in", 6, {"fmt": "0"}), ("Waler Section", "in", 12, {}), ("Strut / Corner Brace", "in", 14, {}),
    ("Base Stability FOS", "in", 7, {"fmt": "0.00"}), ("Dewatering / Grouting", "in", 16, {}),
    ("Max Wall Deflection (mm)", "in", 8, {"fmt": "0"}), *admin_cols("EP"), *file_cols("f_ep"), *qty_cols(),
    ("BCA ERSS Submission (indicative)", "calc", 24, {"formula": f_bca("ERSS Type", EXD, lambda c: "OR(" + ",".join(f'{c("ERSS Type")}="{t}"' for t in NONRC_EP) + ")")}),
    ("SG Compliance Check", "calc", 44, {"formula": f_comp("ERSS Type", ep_parts)}),
    *TAIL, ("Selectable", "calc", 8, {"formula": f_sel("ERSS Type")}),
    ("Selector Key", "calc", 40, {"formula": f_key("ERSS Type", ["ERSS Type", EXD, PL, "Soil Profile"])}),
]
CB_EP = "SS EN 1997-1 DA1 + NA; BCA ERSS requirements (Mar 2024); BCA sheet pile guideline (Nov 2025)"


def ep(t, pur, l, w, d, soil, gwt, q, sec, ln, toe, lv, wal, strut, fos, dew, dfl, rev, status, dwg, scale, staad, qq, lim):
    e = {"ERSS Type": t, "Purpose": pur, PL: l, PW: w, EXD: d, "Soil Profile": soil, "GWT (m below ground)": gwt,
         "Surcharge (kPa)": q, "Sheet Pile / Wall Section": sec, "Pile Length (m)": ln, "Toe Embedment (m)": toe,
         "Strut Levels": lv, "Waler Section": wal, "Strut / Corner Brace": strut, "Base Stability FOS": fos,
         "Dewatering / Grouting": dew, "Max Wall Deflection (mm)": dfl}
    e.update(ex_common(rev, status, dwg, scale, staad, "nr", qq, lim, CB_EP))
    return e


ep_ex = [
    ep("Trench Shoring Box", "Manhole", 3000, 2500, 3.0, FSC, 2.0, 10, "Steel shoring box", None, None, 0, "-", "Box frame", 2.0, "Sump pump", 10, "A", "Approved", "STD-EP-001", 50, False, (.3, 0, 0, 0, 25), "Hired box; quantities exclude hire."),
    ep("Steel Sheet Pile Cofferdam", "Manhole", 4000, 3500, 6.0, KF, 1.0, 20, "FSP III", 12.0, 6.0, 2, "UB 457x191", "Corner brace UB 406", 1.6, "Well points", 30, "A", "Approved", "STD-EP-002", 50, True, (1.5, 0, 9800, 0, 90), "Steel = sheet piles + walers + struts."),
    ep("Steel Sheet Pile Cofferdam", "Jacking pit", 8000, 5000, 7.0, KF, 1.0, 20, "FSP IV", 15.0, 8.0, 2, "UB 610x229", "CHS 508 + thrust wall", 1.5, "Well points + JGP base", 35, "A", "Approved", "STD-EP-003", 50, True, (12.0, 900, 24500, 0, 290), "Includes RC thrust wall."),
    ep("Steel Caisson / Liner Plate", "Manhole", 3000, 3000, 10.0, OA, 3.0, 20, "Liner plate rings", None, None, 0, "Ring stiffeners", "-", 1.8, "Grouting as needed", 15, "A", "Checked", "STD-EP-004", 50, True, (2.0, 0, 7200, 0, 95), "Awaiting QP approval."),
]
wsp, pi = build_register("Temporary ERSS Pit Register (manholes, jacking / receiving pits, trenches)", "ERSS_Pit_Register",
                         "Temporary works for sewer manholes and pits, per nr. BCA submission route shown is indicative: QP to confirm per project.",
                         ep_cols, ep_ex, "tblERSSPit", "ERSS Pit")

# =============================================================== 6. COVERED LINKWAY (LTA SDRE Ch.22)
LKW_, LKH, LKB = "Roof Width, eave to eave (m)", "Clear Headroom (m)", "Bay / Column Spacing (m)"


def lk_id(r, i):
    c = C(i, r)
    return (f'=IF({c("Linkway Type")}="","","LKW-"&{vlk(c("Linkway Type"),"LkType")}&"-W"&TEXT({c(LKW_)}*1000,"0000")&"-H"&TEXT({c(LKH)}*1000,"0000")'
            f'&"-B"&TEXT({c(LKB)}*1000,"0000")&"-"&{c("Rev")})')


def lk_desc(r, i):
    c = C(i, r)
    return (f'=IF({c("Linkway Type")}="","",{c("Linkway Type")}&", "&{c(LKW_)}&" m wide, "&{c(LKH)}&" m headroom, "&{c(LKB)}&" m bays, "'
            f'&{c("Column Arrangement")}&" ("&{c("SDRE Drawing Ref")}&")")')


def lk_parts(c):
    t = c("Linkway Type")
    return [f'IF(AND({t}<>"High Covered Linkway",{c(LKH)}<{S["lkH"]}),"FAIL: headroom below "&{S["lkH"]}&" m. ","")',
            f'IF(AND({t}="High Covered Linkway",{c(LKH)}<{S["lkHH"]}),"FAIL: high linkway headroom below "&{S["lkHH"]}&" m. ","")',
            f'IF({c(LKW_)}<{S["lkW"]},"FAIL: roof width below "&{S["lkW"]}&" m. ",IF({c(LKW_)}<{S["lkW2"]},"NOTE: not for MRT/LRT or high-traffic routes (min "&{S["lkW2"]}&" m). ",""))',
            f'IF(AND({c("Basic Wind Speed (m/s)")}<>"",{c("Basic Wind Speed (m/s)")}<{S["lkV"]}),"FAIL: wind speed below "&{S["lkV"]}&" m/s. ","")',
            f'IF(AND({c("Roof Fixing Pull-out (kN)")}<>"",{c("Roof Fixing Pull-out (kN)")}<{S["lkP"]}),"FAIL: roof fixing pull-out below "&{S["lkP"]}&" kN. ","")',
            grade_chk(c, "Concrete Grade", "lkC"),
            f'IF({c("SDRE Drawing Ref")}="Non-standard (project design)","NOTE: not an SDRE standard, needs LTA acceptance. ","")']


lk_cols = [
    ("Design ID", "calc", 32, {"formula": lk_id}), ("Description", "calc", 54, {"formula": lk_desc}),
    ("Linkway Type", "in", 24, {"dv": "LkType"}), ("SDRE Drawing Ref", "in", 20, {"dv": "LkDwg"}),
    (LKW_, "in", 8, {"fmt": "0.0"}), (LKH, "in", 8, {"fmt": "0.0"}), (LKB, "in", 8, {"fmt": "0.0"}),
    ("Column Arrangement", "in", 18, {"dv": "LkCols"}), ("Column Section", "in", 14, {}), ("Roof Material", "in", 16, {"dv": "Roof"}),
    ("Roof Profile / Fall", "in", 12, {}), ("Gutter & Downpipe", "in", 14, {}), ("Footing Type", "in", 12, {}),
    ("Footing Size (L x B x D mm)", "in", 14, {}), ("RC Stump (dia x H mm)", "in", 12, {}), ("Basic Wind Speed (m/s)", "in", 7, {"fmt": "0"}),
    ("Roof Fixing Pull-out (kN)", "in", 7, {"fmt": "0.0"}), ("Concrete Grade", "in", 8, {"dv": "Grade"}),
    ("Demountable / Bolted (Y/N)", "in", 8, {"dv": "YN"}), ("Lighting Provision", "in", 12, {}),
    ("Roof Sheeting (m2/unit)", "in", 8, {"fmt": "0.0"}), *admin_cols("LK"), *file_cols("f_lk"), *qty_cols(),
    ("SG Compliance Check", "calc", 44, {"formula": f_comp("Linkway Type", lk_parts)}),
    *TAIL, ("Selectable", "calc", 8, {"formula": f_sel("Linkway Type")}),
    ("Selector Key", "calc", 34, {"formula": f_key("Linkway Type", ["Linkway Type", LKW_, LKH])}),
]
CB_LK = "LTA SDRE Rev I (Mar 2026) Ch.22; LTA ADC s4.8; SS EN 1991-1-4; SS EN 1993-1-1"


def lk(t, ref, w, h, b, cols_, csec, roof, prof, fdn, fsize, stump, rev, status, dwg, scale, staad, sheet, qq, lim):
    e = {"Linkway Type": t, "SDRE Drawing Ref": ref, LKW_: w, LKH: h, LKB: b, "Column Arrangement": cols_, "Column Section": csec,
         "Roof Material": roof, "Roof Profile / Fall": prof, "Gutter & Downpipe": "Aluminium gutter, 100 dia DP", "Footing Type": fdn,
         "Footing Size (L x B x D mm)": fsize, "RC Stump (dia x H mm)": stump, "Basic Wind Speed (m/s)": 22,
         "Roof Fixing Pull-out (kN)": 4.0, "Concrete Grade": "C25/30", "Demountable / Bolted (Y/N)": "Y",
         "Lighting Provision": "Conduit only", "Roof Sheeting (m2/unit)": sheet}
    e.update(ex_common(rev, status, dwg, scale, staad, "bay", qq, lim, CB_LK))
    return e


lk_ex = [
    lk("Low Covered Linkway", "LTA/SDRE25/22/CLW1", 2.4, 2.4, 6.0, "Single column, central", "CHS 168.3x6.3", "Aluminium sheet", "Curved", "Pad footing", "1200 x 1200 x 600", "400 x 600", "A", "Approved", "STD-LKW-001", 50, True, 15.6, (1.0, 45, 380, 2.9, 1.4), "Per bay incl. 1 footing. Confirm against SDRE CLW sheets."),
    lk("Low Covered Linkway", "LTA/SDRE25/22/CLW1", 3.6, 2.4, 6.0, "Double column", "CHS 139.7x5.0", "Aluminium + polycarbonate", "Curved", "Pad footing", "1000 x 1000 x 600", "400 x 600", "A", "Approved", "STD-LKW-002", 50, True, 23.0, (1.4, 60, 520, 4.0, 2.0), "For MRT/LRT and high-traffic routes."),
    lk("High Covered Linkway", "LTA/SDRE25/22/CLW2", 3.6, 5.7, 8.0, "Double column", "CHS 219.1x8.0", "Metal deck", "Flat, 2% fall", "Pad footing", "1500 x 1500 x 800", "500 x 600", "A", "Checked", "STD-LKW-003", 100, True, 30.5, (3.6, 120, 1450, 6.2, 4.5), "Over vehicular access. Awaiting QP approval."),
    lk("Modular Covered Linkway (SDRE Rev I)", "LTA/SDRE25/22/CLW3", 2.4, 2.4, 6.0, "Single column, side (cantilever)", "Per SDRE CLW3", "Aluminium sheet", "Per SDRE", "Pad footing", "Per SDRE", "Per SDRE", "A", "WIP", "STD-LKW-004", 50, True, 15.6, (1.1, 50, 420, 3.0, 1.5), "Populate from SDRE Rev I modular sheets (CLW3-CLW5)."),
]
wsl, li = build_register("Covered Linkway Register (LTA SDRE Ch.22)", "Linkway_Register",
                         "Follows LTA SDRE Rev I (Mar 2026) Chapter 22 Covered Linkway (CLW1 layout/elevations, CLW2 high linkway, CLW3-CLW5 modular). Quantities per bay.",
                         lk_cols, lk_ex, "tblLinkway", "Linkway")

# =============================================================== 7. OTHERS
def ot_id(r, i):
    c = C(i, r)
    return (f'=IF({c("Family")}="","",{vlk(c("Family"),"Family")}&"-"&UPPER(SUBSTITUTE({c("Type (short)")}," ",""))'
            f'&IF({c("Span / Length (m)")}<>"","-SP"&TEXT({c("Span / Length (m)")}*1000,"0"),"")'
            f'&IF({c("Height / Rise (m)")}<>"","-H"&TEXT({c("Height / Rise (m)")}*1000,"0"),"")&"-"&{c("Rev")})')


def ot_desc(r, i):
    c = C(i, r)
    return (f'=IF({c("Family")}="","",{c("Family")}&" - "&{c("Type (short)")}&IF({c("Span / Length (m)")}<>"",", span "&{c("Span / Length (m)")}&" m","")'
            f'&IF({c("Height / Rise (m)")}<>"",", H "&{c("Height / Rise (m)")}&" m","")&IF({c("Width (m)")}<>"",", W "&{c("Width (m)")}&" m",""))')


ot_cols = [
    ("Design ID", "calc", 30, {"formula": ot_id}), ("Description", "calc", 40, {"formula": ot_desc}),
    ("Family", "in", 13, {"dv": "Family"}), ("Type (short)", "in", 18, {}), ("Span / Length (m)", "in", 7, {"fmt": "0.0"}),
    ("Height / Rise (m)", "in", 7, {"fmt": "0.0"}), ("Width (m)", "in", 6, {"fmt": "0.0"}), ("Design Load", "in", 18, {}),
    ("Foundation", "in", 14, {}), ("Material", "in", 14, {}), ("Key Members / Reinforcement", "in", 22, {}),
    *admin_cols("OT"), *file_cols("f_ot"), *qty_cols(),
    ("SG Compliance Check", "calc", 20, {"formula": lambda r, i: f'=IF({i["Family"]}{r}="","","OK")'}),
    *TAIL, ("Selectable", "calc", 8, {"formula": f_sel("Family")}),
]
CB_OT = "LTA CDC C&S A3; LTA SDRE Rev I; SS EN 1990/1991/1993"


def ot(fam, typ, sp, h, w, load, fdn, mat, mem, rev, stt, dwg, scale, staad, unit, qq, lim):
    e = {"Family": fam, "Type (short)": typ, "Span / Length (m)": sp, "Height / Rise (m)": h, "Width (m)": w,
         "Design Load": load, "Foundation": fdn, "Material": mat, "Key Members / Reinforcement": mem}
    e.update(ex_common(rev, stt, dwg, scale, staad, unit, qq, lim, CB_OT))
    return e


ot_ex = [
    ot("Railing", "Drain Safety", None, 1.0, None, "Line load per SS EN 1991-1-1", "Cast-in plate", "Galv. steel", "50x50 SHS posts @1.5 m", "A", "Approved", "STD-RLG-001", 10, False, "m", (0, 0, 18, 0, 0), "For open drains deeper than 1.0 m."),
    ot("Railing", "Parapet", None, 1.1, None, "Line load per SS EN 1991-1-1", "Bolted to RC kerb", "Galv. steel", "100x50 RHS posts @1.2 m", "A", "Checked", "STD-RLG-002", 10, False, "m", (0, 0, 32, 0, 0), ""),
    ot("Staircase", "RC Straight Flight", 3.6, 3.0, 1.5, "Imposed per SS EN 1991-1-1", "Ground beam", "RC C32/40", "Waist 175, T13@150", "A", "Approved", "STD-STR-001", 25, False, "nr", (2.4, 210, 0, 12.5, 1.8), "Per flight."),
    ot("POB", "Steel Truss with Lift", 30.0, 5.7, 3.0, "Crowd load; SS EN 1991-2", "Bored piles", "Steel truss", "Warren truss, CHS chords", "A", "WIP", "STD-POB-001", 100, True, "nr", (85, 6000, 46000, 210, 120), "Indicative only."),
]
wso, oi = build_register("Other Structures Register", "Other_Register",
                         "Railing, staircase, POB, USD, bus shelter and others.", ot_cols, ot_ex, "tblOther", "Others")

# =============================================================== helpers over registers
def rng(sheet, idx, h):
    c = idx[h]
    return f"{sheet}!${c}${FIRST}:${c}${LAST}"


DS, RS, ES, MS, PS, LS, OS = "Drain_Register", "RW_Register", "Embedded_Wall_Register", "Manhole_Register", "ERSS_Pit_Register", "Linkway_Register", "Other_Register"
IDX = {s: i for s, i, _ in REGS}


def lk_all(rref, h, default='""'):
    """Look up column h for the Design ID in rref across all registers."""
    expr = default
    for sheet, idx, _lab in reversed(REGS):
        if h in idx:
            expr = f'IFERROR(INDEX({rng(sheet, idx, h)},MATCH({rref},{rng(sheet, idx, "Design ID")},0)),{expr})'
    return expr


def reg_label(rref):
    expr = '""'
    for sheet, idx, lab in reversed(REGS):
        expr = f'IF(ISNUMBER(MATCH({rref},{rng(sheet, idx, "Design ID")},0)),"{lab}",{expr})'
    return expr


# =============================================================== DESIGN SELECTOR
ss = wb.create_sheet("Design_Selector", 1)
ss["A1"] = "Design Selector: nearest APPROVED, SG-compliant standard design"
ss["A1"].font = f_title
ss["A2"] = ("Type into the yellow cells. Each block returns the smallest design at or above the requirement whose Status is Approved and whose SG "
            "compliance check has no FAIL. Linkways and others: filter the register directly (few variants).")
ss["A2"].font = f_sub
for col, w in zip("ABCDEFGHIJK", [32, 46, 2, 32, 46, 2, 32, 46, 2, 32, 46]):
    ss.column_dimensions[col].width = w


def block(lc, vc, title, inputs, steps, rowref_cell, sheet, idx, outs, result):
    put(ss, f"{lc}4", title, f_hdr, fill_hdr["in"])
    put(ss, f"{vc}4", "", f_hdr, fill_hdr["in"])
    for rr, lab, val, dvn in inputs:
        put(ss, f"{lc}{rr}", lab, f_bold)
        put(ss, f"{vc}{rr}", val, f_in, fill_in)
        if dvn:
            dv_list(ss, dvn, f"{vc}{rr}")
    for rr, lab, f, fmt in steps:
        put(ss, f"{lc}{rr}", lab, f_bold)
        put(ss, f"{vc}{rr}", f, f_calc, fill_calc, fmt)
    rref = f"${vc}${rowref_cell}"
    for k, (lab, h) in enumerate(outs):
        r = 18 + k
        put(ss, f"{lc}{r}", lab, f_bold)
        if lab == "Result":
            f = result
        elif callable(h):
            f = h(rref)
        else:
            f = f'=IF({rref}=0,"",INDEX({rng(sheet, idx, h)},{rref}))'
        put(ss, f"{vc}{r}", f, Font(name=F, size=10, bold=lab in ("Result", "Design ID")), fill_calc, al=topwrap)


def common_outs(sheet, idx, extra):
    j = lambda a, b: (lambda rr: f'=IF({rr}=0,"",INDEX({rng(sheet,idx,a)},{rr})&" / "&INDEX({rng(sheet,idx,b)},{rr}))')
    return ([("Result", None), ("Design ID", "Design ID"), ("Description", "Description"), ("Rev / Drawing No.", j("Rev", "Drawing No."))]
            + extra +
            [("SG compliance", "SG Compliance Check"), ("Detail scale 1:", "Detail Scale (1:n)"), ("Design folder", "Design Folder"),
             ("CAD detail", lambda rr: f'=IF({rr}=0,"",INDEX({rng(sheet,idx,"Design Folder")},{rr})&INDEX({rng(sheet,idx,"CAD Detail (.dwg)")},{rr}))'),
             ("STAAD.Pro", "STAAD.Pro (.std)"), ("Calc sheet / report", j("Calc Sheet (.xlsx)", "Calc Report (.pdf)")),
             ("Concrete / Rebar per unit", j("Concrete (m3/unit)", "Rebar (kg/unit)")), ("Limitations", "Limitations / Notes")])


NOMATCH = '"NO STANDARD DESIGN: bespoke design needed"'
# --- Drain
i = di
dT, dW, dD, dLd, dSel, dK, dQ = (rng(DS, i, h) for h in ["Drain Type", WD, DD, "Loading", "Selectable", "Selector Key", "Flow Capacity Q (m3/s)"])
qcrit = 'IF(N(B9)>0,">="&B9,"<>#")'
vel = rng(DS, i, "Velocity at Capacity (m/s)")
block("A", "B", "DRAIN / CULVERT / SUMP",
      [(5, "Drain type", "Precast U-Drain", "DrainType"), (6, "Min. internal width (mm)", 600, None), (7, "Min. internal depth (mm)", 600, None),
       (8, "Loading", S12, "Loading"), (9, "Design flow Q (m3/s), blank = ignore", 0.35, None), (11, "Drain gradient 1 : S  (S =)", 500, None)],
      [(13, "Selected width (mm)", f'=_xlfn.MINIFS({dW},{dT},B5,{dLd},B8,{dSel},"Y",{dW},">="&B6,{dD},">="&B7,{dQ},{qcrit})', None),
       (14, "Selected depth (mm)", f'=IF(B13=0,0,_xlfn.MINIFS({dD},{dT},B5,{dLd},B8,{dSel},"Y",{dW},B13,{dD},">="&B7,{dQ},{qcrit}))', None),
       (15, "Row in register", f'=IFERROR(MATCH(B5&"|"&B13&"|"&B14&"|"&B8&"|Y",{dK},0),0)', None)],
      15, DS, i,
      common_outs(DS, i, [("Capacity Q / V at capacity", lambda rr: f'=IF({rr}=0,"",INDEX({dQ},{rr})&" m3/s / "&INDEX({vel},{rr})&" m/s")'),
                          ("Velocity check", lambda rr: f'=IF({rr}=0,"",IF(INDEX({vel},{rr})="","n/a",IF(INDEX({vel},{rr})>{S["vmax"]},"NOTE: V > "&{S["vmax"]}&" m/s, check scour",IF(INDEX({vel},{rr})<{S["vmin"]},"NOTE: V < "&{S["vmin"]}&" m/s","OK"))))')]),
      f'=IF(B15=0,{NOMATCH},IF(OR(B13>B6,B14>B7),"Match (conservative: larger than required)","Exact match"))')
ss["A10"] = "Capacity uses Manning n and freeboard from SG_Criteria."
ss["A10"].font = f_sub
# --- RC wall
i = ri
rT, rH, rQ, rG, rSel, rK = (rng(RS, i, h) for h in ["Wall Type", RWH, "Surcharge (kPa)", "GWT Behind Wall (Y/N)", "Selectable", "Selector Key"])
gcrit = 'IF(E8="Y","Y","<>#")'
block("D", "E", "RC / GRAVITY RETAINING WALL",
      [(5, "Wall type", CANT, "WallType"), (6, "Retained height (m)", 2.8, None), (7, "Surcharge (kPa)", 25, None), (8, "GWT behind wall? (Y/N)", "N", "YN")],
      [(13, "Selected height (m)", f'=_xlfn.MINIFS({rH},{rT},E5,{rSel},"Y",{rH},">="&E6,{rQ},">="&E7,{rG},{gcrit})', "0.00"),
       (14, "Selected surcharge (kPa)", f'=IF(E13=0,0,_xlfn.MINIFS({rQ},{rT},E5,{rSel},"Y",{rH},E13,{rQ},">="&E7,{rG},{gcrit}))', None),
       (15, "Selected GWT case", f'=IF(E13=0,"",IF(E8="Y","Y",IF(COUNTIFS({rT},E5,{rSel},"Y",{rH},E13,{rQ},E14,{rG},"N")>0,"N","Y")))', None),
       (16, "Row in register", f'=IFERROR(MATCH(E5&"|"&E13&"|"&E14&"|"&E15&"|Y",{rK},0),0)', None)],
      16, RS, i,
      common_outs(RS, i, [("Stem / base bars", lambda rr: f'=IF({rr}=0,"",INDEX({rng(RS,ri,"Stem Main Bar")},{rr})&" / "&INDEX({rng(RS,ri,"Base Main Bar")},{rr}))'),
                          ("BCA ERSS submission", "BCA ERSS Submission (indicative)")]),
      f'=IF(E16=0,{NOMATCH},IF(OR(E13>E6,E14>E7,E15<>E8),"Match (conservative: covers a worse case)","Exact match"))')
# --- Embedded wall
i = ei
eT, eD, eSu, eSo, eTP, eSel, eK = (rng(ES, i, h) for h in ["Wall Type", EXD, "Support System", "Soil Profile", "Temporary / Permanent", "Selectable", "Selector Key"])
block("G", "H", "EMBEDDED WALL (CBP / SECANT / SHEET PILE / D-WALL)",
      [(5, "Wall type", "Contiguous Bored Pile (CBP)", "EmbType"), (6, "Excavation depth (m)", 4.5, None),
       (7, "Support system", "1 level strut", "Support"), (8, "Soil profile", OA, "Soil"), (9, "Temporary / Permanent", "Permanent", "TempPerm")],
      [(13, "Selected depth (m)", f'=_xlfn.MINIFS({eD},{eT},H5,{eSu},H7,{eSo},H8,{eTP},H9,{eSel},"Y",{eD},">="&H6)', "0.0"),
       (15, "Row in register", f'=IFERROR(MATCH(H5&"|"&H13&"|"&H7&"|"&H8&"|"&H9&"|Y",{eK},0),0)', None)],
      15, ES, i,
      common_outs(ES, i, [("Pile / length / toe", lambda rr: f'=IF({rr}=0,"",INDEX({rng(ES,ei,"Pile Dia / Section")},{rr})&", L = "&INDEX({rng(ES,ei,"Pile / Wall Length (m)")},{rr})&" m, toe "&INDEX({rng(ES,ei,"Toe Embedment (m)")},{rr})&" m")'),
                          ("BCA ERSS submission", "BCA ERSS Submission (indicative)")]),
      f'=IF(H15=0,{NOMATCH},IF(H13>H6,"Match (conservative: deeper than required)","Exact match"))')
# --- Manhole
i = mi
mT, mDx, mDn, mL, mP, mSel, mK = (rng(MS, i, h) for h in ["Manhole Type", DMAX, DMIN, MHL, PIPE, "Selectable", "Selector Key"])
block("J", "K", "SEWER MANHOLE",
      [(5, "Manhole type", SMH, "MHType"), (6, "Depth to invert (m)", 4.2, None), (7, "Largest incoming pipe dia (mm)", 450, None)],
      [(13, "Selected max depth (m)", f'=_xlfn.MINIFS({mDx},{mT},K5,{mSel},"Y",{mDx},">="&K6,{mDn},"<="&K6,{mP},">="&K7)', "0.0"),
       (14, "Selected internal length (mm)", f'=IF(K13=0,0,_xlfn.MINIFS({mL},{mT},K5,{mSel},"Y",{mDx},K13,{mDn},"<="&K6,{mP},">="&K7))', None),
       (15, "Row in register", f'=IFERROR(MATCH(K5&"|"&K13&"|"&K14&"|Y",{mK},0),0)', None)],
      15, MS, i,
      common_outs(MS, i, [("PUB standard drawing", "PUB Std Drawing Ref"),
                          ("Typical temporary ERSS", "Typical Temporary ERSS (Design ID)")]),
      f'=IF(K15=0,{NOMATCH},"Match: depth range "&INDEX({mDn},K15)&" - "&INDEX({mDx},K15)&" m")')
ss["A36"] = ("Loading matched exactly. A wall designed with groundwater is offered for a no-groundwater search, not the reverse. Velocity is at full capacity. "
             "BCA submission route is indicative only; the QP confirms. Copy the Design ID into Tender_Pack.")
ss["A36"].font = f_sub

# =============================================================== TENDER PACK
tp = wb.create_sheet("Tender_Pack", 2)
tp["A1"] = "Tender Pack: designs used in this tender, quantities, and CAD sheet layout"
tp["A1"].font = f_title
tp["A2"] = "Type Design IDs from any register, length / number, and the CAD sheet each detail goes on. Then save CAD_Export for AutoCAD. One copy of this workbook per tender."
tp["A2"].font = f_sub
put(tp, "A3", "Tender / Project:", f_bold)
put(tp, "B3", "EXAMPLE - Tender T-2026-01 Road and Sewer Upgrading at Jalan Contoh", f_in, fill_in)
tp.merge_cells("B3:D3")
tph = [("Item", 5, "in"), ("Design ID", 34, "in"), ("Location / Remarks", 24, "in"), ("Length / Qty", 8, "in"),
       ("CAD Sheet No.", 7, "in"), ("Unit", 5, "calc"), ("Register", 11, "calc"), ("Description", 48, "calc"),
       ("Rev", 5, "calc"), ("Status", 9, "calc"), ("Drawing No.", 12, "calc"), ("Detail Scale 1:", 7, "calc"),
       ("CAD Detail File (full path)", 60, "calc"), ("Concrete (m3)", 9, "qty"), ("Rebar (t)", 8, "qty"), ("Steel (t)", 8, "qty"),
       ("Formwork (m2)", 9, "qty"), ("Excavation (m3)", 9, "qty"), ("Check", 34, "calc")]
TH, T1, TN = 5, 6, 60
TC = {h: L(c) for c, (h, *_r) in enumerate(tph, 1)}
for c, (h, w, k) in enumerate(tph, 1):
    cell = tp.cell(TH, c, h)
    cell.font, cell.alignment, cell.border, cell.fill = f_hdr, wrap, box, fill_hdr[k]
    tp.column_dimensions[L(c)].width = w
tp.row_dimensions[TH].height = 40
for r in range(T1, TN + 1):
    for c, (h, w, k) in enumerate(tph, 1):
        tp.cell(r, c).border = box
        tp.cell(r, c).font = f_in if k == "in" else f_calc
        tp.cell(r, c).fill = fill_in if k == "in" else fill_calc
    b = f"$B{r}"
    g = f'IF({b}="","",'
    tp[f"{TC['Register']}{r}"] = f"={g}{reg_label(b)})"
    tp[f"{TC['Unit']}{r}"] = f"={g}{lk_all(b, 'Qty Unit')})"
    tp[f"{TC['Description']}{r}"] = f'={g}{lk_all(b, "Description", chr(34)+"ID not found - check spelling"+chr(34))})'
    for h, src in (("Rev", "Rev"), ("Status", "Status"), ("Drawing No.", "Drawing No."), ("Detail Scale 1:", "Detail Scale (1:n)")):
        tp[f"{TC[h]}{r}"] = f"={g}{lk_all(b, src)})"
    tp[f"{TC['CAD Detail File (full path)']}{r}"] = f"={g}{lk_all(b, 'Design Folder')}&{lk_all(b, 'CAD Detail (.dwg)')})"
    for h, src, div in (("Concrete (m3)", QTY[0], ""), ("Rebar (t)", QTY[1], "/1000"), ("Steel (t)", QTY[2], "/1000"),
                        ("Formwork (m2)", QTY[3], ""), ("Excavation (m3)", QTY[4], "")):
        tp[f"{TC[h]}{r}"] = f'=IF(OR({b}="",$D{r}=""),"",$D{r}*N({lk_all(b, src, "0")}){div})'
        tp[f"{TC[h]}{r}"].number_format = "#,##0.00" if "(t)" in h else "#,##0.0"
    reg, stc = TC["Register"], TC["Status"]
    tp[f"{TC['Check']}{r}"] = (f'={g}IF({reg}{r}="","ID not found",IF({stc}{r}<>"Approved","WARNING: status is "&{stc}{r},'
                               f'IF({lk_all(b, "Selectable")}="N","WARNING: fails SG compliance check",'
                               f'IF(D{r}="","Enter length / qty",IF(E{r}="","Enter CAD sheet no.","OK"))))))')
tr = TN + 1
put(tp, f"{TC['CAD Detail File (full path)']}{tr}", "TOTAL", f_bold)
for h in ("Concrete (m3)", "Rebar (t)", "Steel (t)", "Formwork (m2)", "Excavation (m3)"):
    c = TC[h]
    put(tp, f"{c}{tr}", f"=SUM({c}{T1}:{c}{TN})", f_bold, fill_calc, "#,##0.0")
tp[f"A{tr+2}"] = "Quantities are PRELIMINARY (standard section x length or per nr) for tender estimation only. QS to verify."
tp[f"A{tr+2}"].font = f_sub
tex = [(1, "DRN-PUD-0600x0600-S12-A", "Roadside drain, Ch 0-250 both sides", 500, 1),
       (2, "DRN-SUD-0600x0750-FWY-A", "Under new footway", 120, 1),
       (3, "DRN-BC-1500x1200-VEH-A", "Road crossing at Ch 180", 18, 2),
       (4, "DRN-SP-0900x1200-S12-A", "Drain sumps", 13, 2),
       (5, "RW-CANT-H3000-S25-A", "Slope at carpark, grid A-D", 42, 3),
       (6, "EW-CBP-D5000-S1-OA-P-A", "Basement edge to adjacent road", 35, 3),
       (7, "MH-SM-1500x1200-D6000-C-A", "Sewer MH1-MH4", 4, 4),
       (8, "ERSS-SSP-4000x3500-D6000-KF-A", "Temporary ERSS for MH1-MH4", 4, 4),
       (9, "LKW-LCL-W3600-H2400-B6000-A", "Linkway to MRT exit", 12, 5),
       (10, "DRN-OD-2000x1800-S12-A", "Outlet drain to canal", 60, 2)]
for k, row in enumerate(tex):
    rr = T1 + k
    for c, v in zip("ABCDE", row):
        tp[f"{c}{rr}"] = v
        tp[f"{c}{rr}"].fill = fill_ex
tp.freeze_panes = "C6"

# =============================================================== CAD EXPORT
cx = wb.create_sheet("CAD_Export", 3)
cx["A1"], cx["B1"] = "PROJECT", "=Tender_Pack!B3"
cx["A2"], cx["B2"] = "DWGPREFIX", "T2026-01-STD"
cx["A3"], cx["B3"] = "LIBROOT", f"={S['rootn']}"
for a in ("A1", "A2", "A3"):
    cx[a].font = f_bold
cx["B1"].font, cx["B3"].font = f_calc, f_calc
cx["B2"].font, cx["B2"].fill = f_in, fill_in
cx["D1"] = "HOW TO USE: (1) fill Tender_Pack incl. CAD Sheet No.  (2) with THIS sheet active: File > Save As > Text (Tab delimited) (*.txt)"
cx["D2"] = "(3) AutoCAD: new drawing from your title-block template, APPLOAD TenderSheets.lsp, run TPSHEETS, pick the .txt"
cx["D3"] = "Only Approved rows are placed. Keep rows 1-5 and the column order: the AutoCAD routine reads them."
for a in ("D1", "D2", "D3"):
    cx[a].font = f_sub
for c, h in enumerate(["SHEET", "DETAIL", "CADFILE", "SCALE", "TITLE", "DESIGNID", "REV", "STATUS"], 1):
    put(cx, f"{L(c)}5", h, f_hdr, fill_hdr["calc"])
for k in range(TN - T1 + 1):
    r, t_ = 6 + k, T1 + k
    b = f"Tender_Pack!$B${t_}"
    cx[f"A{r}"] = f'=IF({b}="","",IF(Tender_Pack!$E${t_}="",1,Tender_Pack!$E${t_}))'
    cx[f"B{r}"] = f'=IF(OR(A{r}="",H{r}<>"Approved"),"",COUNTIFS($A$6:A{r},A{r},$H$6:H{r},"Approved"))'
    cx[f"C{r}"] = f'=IF(A{r}="","",Tender_Pack!${TC["CAD Detail File (full path)"]}${t_})'
    cx[f"D{r}"] = f'=IF(A{r}="","",Tender_Pack!${TC["Detail Scale 1:"]}${t_})'
    cx[f"E{r}"] = f'=IF(A{r}="","",UPPER(Tender_Pack!${TC["Description"]}${t_}))'
    cx[f"F{r}"] = f'=IF(A{r}="","",{b})'
    cx[f"G{r}"] = f'=IF(A{r}="","",Tender_Pack!${TC["Rev"]}${t_})'
    cx[f"H{r}"] = f'=IF(A{r}="","",Tender_Pack!${TC["Status"]}${t_})'
    for c in range(1, 9):
        cx.cell(r, c).font = f_calc
for col, w in zip("ABCDEFGH", [10, 8, 70, 7, 60, 34, 5, 10]):
    cx.column_dimensions[col].width = w

# =============================================================== COVERAGE
cv = wb.create_sheet("Coverage", 4)
cv["A1"] = "Coverage Matrix: number of SELECTABLE designs (Approved, no FAIL). Red 0 = gap"
cv["A1"].font = f_title
cv.column_dimensions["A"].width = 44
cr = 3


def matrix(title, rows_, sheet, idx, rowh, colh, bands=None, exact=None, label="Type"):
    global cr
    cv[f"A{cr}"] = title
    cv[f"A{cr}"].font = f_bold
    cr += 1
    heads = [f"{bands[j-1]:g}-{bands[j]:g}" for j in range(1, len(bands))] if bands else exact
    put(cv, f"A{cr}", label, f_hdr, fill_hdr["in"])
    for j, h in enumerate(heads):
        put(cv, f"{L(j+2)}{cr}", h, f_hdr, fill_hdr["in"], al=wrap)
        cv.column_dimensions[L(j + 2)].width = 9
    hr = cr
    for t in rows_:
        cr += 1
        put(cv, f"A{cr}", t)
        for j in range(len(heads)):
            if bands:
                f = (f'=COUNTIFS({rng(sheet,idx,rowh)},$A{cr},{rng(sheet,idx,"Selectable")},"Y",'
                     f'{rng(sheet,idx,colh)},">{bands[j]}",{rng(sheet,idx,colh)},"<={bands[j+1]}")')
            else:
                f = f'=COUNTIFS({rng(sheet,idx,rowh)},$A{cr},{rng(sheet,idx,"Selectable")},"Y",{rng(sheet,idx,colh)},{L(j+2)}${hr})'
            put(cv, f"{L(j+2)}{cr}", f, f_calc, al=Alignment(horizontal="center"))
    rg = f"B{hr+1}:{L(len(heads)+1)}{cr}"
    cv.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=["0"], fill=fill_red))
    cv.conditional_formatting.add(rg, CellIsRule(operator="greaterThan", formula=["0"], fill=fill_grn))
    cr += 2


names = lambda k: [v for v, _ in lists[k]]
matrix("Drains & culverts: type vs internal width (mm)", names("DrainType"), DS, di, "Drain Type", WD,
       exact=[450, 600, 750, 900, 1050, 1200, 1500, 1800, 2000, 2400, 3000], label="Drain type")
matrix("RC / gravity retaining walls: type vs retained height band (m)", names("WallType"), RS, ri, "Wall Type", RWH,
       bands=[0, 1, 1.5, 2, 2.5, 3, 3.5, 4, 4.5, 5, 6, 8], label="Wall type")
matrix("Embedded walls: type vs excavation depth band (m)", names("EmbType"), ES, ei, "Wall Type", EXD,
       bands=[0, 3, 4.5, 6, 9, 12, 15, 20, 30], label="Wall type")
matrix("Sewer manholes: type vs max depth band (m)", names("MHType"), MS, mi, "Manhole Type", DMAX,
       bands=[0, 3, 4.5, 6, 8, 10, 15], label="Manhole type")
matrix("Temporary ERSS pits: type vs excavation depth band (m)", names("ERSSType"), PS, pi, "ERSS Type", EXD,
       bands=[0, 1.5, 3, 4, 6, 8, 10, 15], label="ERSS type")
matrix("Covered linkways: type vs roof width (m)", names("LkType"), LS, li, "Linkway Type", LKW_, exact=[2.4, 3.0, 3.6, 4.8], label="Linkway type")
cv[f"A{cr}"] = "Library status by register"
cv[f"A{cr}"].font = f_bold
cr += 1
put(cv, f"A{cr}", "Register", f_hdr, fill_hdr["in"])
heads = ["Approved", "Checked", "WIP", "Superseded", "Approved but FAIL", "Files not checked"]
for j, h in enumerate(heads):
    put(cv, f"{L(j+2)}{cr}", h, f_hdr, fill_hdr["in"], al=wrap)
for sheet, idx, lab in REGS:
    cr += 1
    put(cv, f"A{cr}", lab)
    for j, h in enumerate(heads[:4]):
        put(cv, f"{L(j+2)}{cr}", f'=COUNTIF({rng(sheet,idx,"Status")},"{h}")', f_calc)
    put(cv, f"F{cr}", f'=COUNTIFS({rng(sheet,idx,"Status")},"Approved",{rng(sheet,idx,"Selectable")},"N")', f_calc)
    put(cv, f"G{cr}", f'=COUNTIFS({rng(sheet,idx,"Design ID")},"?*",{rng(sheet,idx,"Files Checked in Folder (Y/N)")},"<>Y")', f_calc)

# =============================================================== CAD STANDARD
cs = wb.create_sheet("CAD_Standard", 5)
cs["A1"] = "CAD Detail Standard: rules every library detail DWG must follow"
cs["A1"].font = f_title
cs["A2"] = "The TPSHEETS AutoCAD routine relies on these rules to drop details into the title block without manual fixing."
cs["A2"].font = f_sub
rules = [
    ("FILE SET PER DESIGN", None),
    ("Folder", r"<Library root>\<register folder>\<Design ID>\    e.g. ...\04_Sewer_Manhole\MH-SM-1500x1200-D6000-C-A\ "),
    ("Files", "<ID>_CAD.dwg (tender detail), <ID>_DWG.pdf (plotted detail), <ID>_STAAD.std (STAAD.Pro input, if used), <ID>_CALC.xlsx (live calc), <ID>_CALC.pdf (signed calc report)"),
    ("STAAD", "Keep the .std input. Put the STAAD output summary into the calc report PDF and state the STAAD.Pro version. Output files (.anl etc.) optional."),
    ("Superseded", r"Move old revisions to ...\99_Superseded\<old ID>\. Never delete."),
    ("DRAWING THE DETAIL (<ID>_CAD.dwg)", None),
    ("Units / scale", "Model space, 1 unit = 1 mm, drawn full size."),
    ("Origin", "Bottom-left of the detail frame at 0,0. Nothing at negative coordinates."),
    ("Detail frame", "Fit inside one SLOT of the tender sheet at the detail scale. Default A1 slot = 190 x 252 mm paper, minus 20 mm title band: at 1:20 the model frame is 3800 x 4640 mm; 1:25 = 4750 x 5800; 1:50 = 9500 x 11600; 1:100 = 19000 x 23200."),
    ("Text & dims", "Size text, dims and hatch for the detail's own scale (2.5 mm text at 1:20 = 50 mm high; DIMSCALE 20). The routine scales the block by 1/n."),
    ("No title", "Do not draw the detail title, scale or ID in the DWG. The routine writes them from the register."),
    ("Clean file", "Bind or remove xrefs (eTransmit), no images, PURGE and AUDIT before saving. Layers per company CAD standard."),
    ("SDRE / PUB based details", "Where a detail follows an LTA SDRE or PUB standard drawing, note the standard drawing number in the detail notes and in the register."),
    ("TITLE BLOCK TEMPLATE", None),
    ("Template", "Company A1 title block DWT, title block as a block with attributes on a paper-space layout named TEMPLATE, inserted at 0,0."),
    ("Attributes", "Default tags: PROJECT, DWG_TITLE, DWG_NO, SCALE, SHEET, DATE. Change the tag names at the top of TenderSheets.lsp to match yours."),
    ("Drawing area", "Default x 25 to 645, y 25 to 569 mm (title block on the right). 3 columns x 2 rows = 6 details per A1 sheet. Change in TenderSheets.lsp."),
]
for k, (a, b) in enumerate(rules, 4):
    cs[f"A{k}"] = a
    if b is None:
        cs[f"A{k}"].font = f_sec
    else:
        cs[f"B{k}"] = b
        cs[f"A{k}"].font, cs[f"B{k}"].font = f_bold, f_body
        cs[f"B{k}"].alignment = topwrap
cs.column_dimensions["A"].width = 24
cs.column_dimensions["B"].width = 120

# =============================================================== README
rd = ws_readme
rd.column_dimensions["A"].width = 28
rd.column_dimensions["B"].width = 112
rows = [
    ("Standard Design Library: Register (Singapore)", None),
    ("Purpose", "Call out approved standard designs that meet Singapore requirements (PUB, LTA, BCA) for tender drawings: pick the design, get its CAD, STAAD.Pro, Excel and PDF files, drop the details into the title block, and get preliminary quantities."),
    ("", ""),
    ("WORKFLOW", None),
    ("1. Select", "Design_Selector: drain / culvert / sump, RC wall, embedded wall, sewer manhole. Returns the nearest approved, SG-compliant design and its files."),
    ("2. Tender pack", "Tender_Pack: Design IDs from any register, lengths / numbers, CAD sheet number. Preliminary quantities."),
    ("3. Export", "CAD_Export: save that sheet as Text (Tab delimited)."),
    ("4. CAD", "AutoCAD: new drawing from the title-block template, APPLOAD TenderSheets.lsp, TPSHEETS, pick the .txt."),
    ("", ""),
    ("REGISTERS (how each is tabulated)", None),
    ("Drain_Register", "Drains, culverts, sumps, chambers. Key: type, internal W x D, cells, loading. Checks: PUB roadside / common drain minimums, railing depth, grade, cover. Flow capacity by Manning."),
    ("RW_Register", "RC and gravity walls. Key: wall type (cantilever with toe & heel / L-wall heel only / reverse L toe only / counterfort / gravity / block / gabion / MSE), retained height, surcharge, GWT. Checks: grade, cover, sliding, toe/heel vs type, LTA surcharge; BCA ERSS route."),
    ("Embedded_Wall_Register", "CBP, secant, sheet pile, diaphragm wall, soldier pile, king post. Key: wall type, excavation depth, support system, soil profile (Kallang, FCBB, Jurong, Old Alluvium, Bukit Timah), temporary / permanent. Checks: deflection, CBP below GWT in marine clay, permanent sheet pile corrosion; BCA ERSS route."),
    ("Manhole_Register", "Permanent sewer manholes per PUB COP SSW 3rd Ed and PUB/WRN/STD 001C / 003B / 004C / 006C. Key: type, internal size, depth range, max incoming pipe. Checks: min depth, 001C depth limit, 004C pipe size, uplift FOS. Links its typical temporary ERSS."),
    ("ERSS_Pit_Register", "Temporary ERSS for manholes, jacking / receiving pits and trenches (sheet pile cofferdam, shoring box, caisson, secant / CBP shaft, open cut). Key: type, pit plan size, depth, soil. Checks: base stability, BCA route."),
    ("Linkway_Register", "Covered linkways per LTA SDRE Rev I (Mar 2026) Ch.22 (CLW1-CLW5). Key: type, roof width, headroom, bay. Checks: headroom 2.4 / 5.7 m, width 2.4 / 3.6 m, wind 22 m/s, fixing pull-out 4.0 kN, C25/30."),
    ("Other_Register", "Railing, staircase, POB, USD, bus shelter, others."),
    ("SG_Criteria", "All Singapore criteria and folder paths used by the checks. Rows marked VERIFY must be confirmed against the current documents before go-live."),
    ("", ""),
    ("COLOUR LEGEND", None),
    ("Yellow cell, blue text", "Input."), ("Grey cell", "Formula: do not overwrite."), ("Purple header", "File set for the design."),
    ("Green header", "Tender quantity per unit."), ("Green-shaded rows", "EXAMPLE data to show the format. Values are illustrative, not designs."),
    ("Compliance cell", "Green OK, amber NOTE (allowed, check), red FAIL (hidden from the Selector)."),
    ("", ""),
    ("SHAREPOINT", None),
    ("Tables", "Each register is an Excel Table (tblDrain, tblRW, tblEmbedded, tblManhole, tblERSSPit, tblLinkway, tblOther) for import to SharePoint Lists / Power Apps."),
    ("Sync", "Sync the SharePoint library to each user's PC and set SG_Criteria 'Library root folder' to the synced path."),
    ("", ""),
    ("DISCLAIMER", "Standard designs are valid only within their stated assumptions. The QP must verify site conditions, loading, soil and authority requirements for every project. ERSS and embedded walls always need site-specific design and BCA submission where required. SG criteria here must be checked against current PUB / LTA / BCA documents. Quantities are preliminary."),
]
for k, (a, b) in enumerate(rows, 1):
    ca = rd.cell(k, 1, a)
    if k == 1:
        ca.font = f_title
    elif b is None:
        ca.font = f_sec
    else:
        ca.font = f_bold
        cb = rd.cell(k, 2, b)
        cb.font, cb.alignment = f_body, topwrap
    ca.alignment = Alignment(vertical="top")

tabc = {"README": "1F3864", "Design_Selector": "C00000", "Tender_Pack": "C00000", "CAD_Export": "C00000",
        "Coverage": "C00000", "CAD_Standard": "7030A0", "SG_Criteria": "BF8F00", "Lists": "7F7F7F"}
for sh in wb.worksheets:
    sh.sheet_properties.tabColor = tabc.get(sh.title, "375623")
wb.save(OUT)
print("saved")
