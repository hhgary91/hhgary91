"""Builds design-library-search.html from Design_Library_Register.xlsx (cached values)."""
import json, datetime, os
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', 'Design_Library_Register.xlsx')
OUT = os.path.join(HERE, '..', 'design-library-search.html')
wb = load_workbook(SRC, data_only=True)


def val(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime('%Y-%m-%d')
    return v


def table(sh, key):
    ws = wb[sh]
    hdr = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
    out = []
    for r in range(5, ws.max_row + 1):
        row = {h: val(ws.cell(r, c + 1).value) for c, h in enumerate(hdr) if h and ws.cell(r, c + 1).value not in (None, '')}
        if row.get(key):
            out.append(row)
    return out


reg = {n: [{k: v for k, v in r.items() if k != 'Selector Key'} for r in table(n, 'Design ID')] for n in wb.sheetnames if n.endswith('_Register')}
data = {'reg': reg, 'dict': table('Data_Dictionary', 'Field'), 'projects': table('Projects', 'Project Ref'),
        'rebar': table('Rebar_Schedule', 'Design ID'), 'revisions': table('Revision_Log', 'Design ID (new rev)'),
        'source': 'Design_Library_Register.xlsx (example data)'}
html = open(os.path.join(HERE, 'tpl_search.html')).read().replace('/*DATA*/', json.dumps(data, ensure_ascii=False))
open(OUT, 'w').write(html)
print(OUT, len(html), {k: len(v) for k, v in reg.items()})
